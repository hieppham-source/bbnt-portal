"""Đọc Google Sheets bằng OAuth2 access token của người dùng: file order (link do user dán) và Sheet master."""
from __future__ import annotations

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any, Protocol
from urllib.parse import parse_qs, quote, urlparse

_SHEET_PATH = re.compile(r"^/spreadsheets/d/([A-Za-z0-9_-]{20,})")
_DRIVE_PATH = re.compile(r"^/file/d/([A-Za-z0-9_-]{20,})")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
#: File Excel tải từ Drive tối đa (byte) — file order vài nghìn dòng thường < 5 MB.
MAX_XLSX_BYTES = 30 * 1024 * 1024
#: Số dòng đầu mỗi tab đọc để dò header khi chọn tab.
HEAD_ROWS = 15


class SheetAccessError(RuntimeError):
    """Không đọc được sheet (chưa có quyền xem, sai gid, Google lỗi...)."""


class SheetAuthError(SheetAccessError):
    """Access token Google của người dùng hết hạn / bị thu hồi -> cần đăng nhập lại."""


class DriveScopeError(SheetAccessError):
    """File là Excel trên Drive nhưng người dùng chưa cấp quyền đọc Drive (scope ``drive.readonly``)."""


@dataclass(frozen=True)
class SheetRef:
    spreadsheet_id: str
    gid: int | None = None

    @property
    def canonical_url(self) -> str:
        url = f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/edit"
        return f"{url}#gid={self.gid}" if self.gid is not None else url


def parse_sheet_url(url: str) -> SheetRef:
    """Nhận link Google Sheets (``/edit#gid=..``, ``/edit?gid=..``, ``/view``...) hoặc link file Drive
    (``drive.google.com/file/d/<id>/view``, ``drive.google.com/open?id=<id>`` — dùng cho file .xlsx).
    Raise ValueError nếu sai."""
    u = urlparse(url.strip())
    q, frag = parse_qs(u.query), parse_qs(u.fragment)
    if u.scheme == "https" and u.netloc == "drive.google.com":
        m = _DRIVE_PATH.match(u.path)
        file_id = m.group(1) if m else (q.get("id", [""])[0] if u.path in ("/open", "/uc") else "")
        if not re.fullmatch(r"[A-Za-z0-9_-]{20,}", file_id or ""):
            raise ValueError("Không tìm thấy id file trong link Drive")
        return SheetRef(file_id)
    if u.scheme != "https" or u.netloc != "docs.google.com":
        raise ValueError("Link phải là https://docs.google.com/spreadsheets/d/... hoặc link file trên Google Drive")
    m = _SHEET_PATH.match(u.path)
    if not m:
        raise ValueError("Không tìm thấy spreadsheet id trong link")
    gid_raw = q.get("gid", [None])[0] or frag.get("gid", [None])[0]
    if gid_raw is not None and not str(gid_raw).isdigit():
        raise ValueError("gid không hợp lệ")
    return SheetRef(m.group(1), int(gid_raw) if gid_raw is not None else None)


@dataclass
class OrderWorkbook:
    """File order đã mở: danh sách tab + cách đọc. Google Sheet đọc qua Sheets API, Excel (.xlsx) trên Drive
    tải về 1 lần rồi đọc bằng openpyxl. Kết quả đọc được nhớ trong request (không lưu ra ngoài)."""

    kind: str                                   # "google_sheet" | "xlsx"
    tabs: list[str]
    reader: Callable[[str], list[list[Any]]]
    head_reader: Callable[[list[str]], dict[str, list[list[Any]]]] | None = None
    gid_tab: str | None = None
    _rows: dict[str, list[list[Any]]] = field(default_factory=dict)

    def read(self, tab: str) -> list[list[Any]]:
        if tab not in self._rows:
            self._rows[tab] = self.reader(tab)
        return self._rows[tab]

    def heads(self, tabs: list[str]) -> dict[str, list[list[Any]]]:
        """HEAD_ROWS dòng đầu của từng tab (để dò header) — 1 request cho Google Sheet."""
        if self.head_reader is not None:
            missing = [t for t in tabs if t not in self._rows]
            got = self.head_reader(missing) if missing else {}
            return {t: (self._rows[t][:HEAD_ROWS] if t in self._rows else got.get(t, [])) for t in tabs}
        return {t: self.read(t)[:HEAD_ROWS] for t in tabs}


def _trim_rows(rows: list[list[Any]]) -> list[list[Any]]:
    while rows and all(c is None or (isinstance(c, str) and not c.strip()) for c in rows[-1]):
        rows.pop()      # ô đã định dạng sẵn nhưng trống -> openpyxl vẫn trả về dòng rỗng
    return rows


def xlsx_workbook(data: bytes) -> OrderWorkbook:
    """Mở file .xlsx (bytes) thành :class:`OrderWorkbook` — chỉ lấy tab đang hiện, giá trị đã tính sẵn."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — zip hỏng, file .xls cũ đổi đuôi...
        raise SheetAccessError("Không mở được file Excel — lưu lại dạng .xlsx hoặc chuyển sang Google Sheet") from exc
    tabs = [ws.title for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]

    def read(tab: str) -> list[list[Any]]:
        return _trim_rows([list(r) for r in wb[tab].iter_rows(values_only=True)])

    def heads(tabs_: list[str]) -> dict[str, list[list[Any]]]:
        return {t: [list(r) for r in wb[t].iter_rows(max_row=HEAD_ROWS, values_only=True)] for t in tabs_}

    return OrderWorkbook("xlsx", tabs, read, heads)


class SheetReader(Protocol):
    def read(self, ref: SheetRef, *, sheet_name: str | None = None) -> list[list[Any]]:
        """Trả về ma trận ô (UNFORMATTED_VALUE, ngày dạng serial number)."""
        ...

    def open(self, ref: SheetRef) -> OrderWorkbook:
        """Mở file order (Google Sheet hoặc .xlsx trên Drive) để chọn tab."""
        ...


class GoogleSheetsReader:
    """Google Sheets API v4 bằng **OAuth2 access token của người dùng** (scope ``spreadsheets.readonly``).

    Frontend Next.js xin quyền này lúc đăng nhập Google và gửi token kèm mỗi request (``X-Google-Access-Token``).
    Ai mở được Sheet nào trên Drive thì tool đọc được Sheet đó cho người đó — không dùng service account.
    Token chỉ sống trong request, không lưu, không ghi log.

    Ưu tiên tab theo ``gid`` trong link; không có gid thì dùng ``sheet_name`` của mapping; không có nữa thì tab đầu.
    """

    BASE = "https://sheets.googleapis.com/v4/spreadsheets"
    DRIVE = "https://www.googleapis.com/drive/v3/files"

    def __init__(self, access_token: str, *, timeout: float = 30.0, session: Any = None) -> None:
        if not access_token:
            raise ValueError("Thiếu access token Google")
        if session is None:
            import requests

            session = requests.Session()
        session.headers["Authorization"] = f"Bearer {access_token}"
        self._session: Any = session
        self._timeout = timeout

    @classmethod
    def from_access_token(cls, token: str, **kw: Any) -> GoogleSheetsReader:
        return cls(token, **kw)

    def _check(self, resp: Any, what: str = "sheet") -> None:
        if resp.status_code == 401:
            raise SheetAuthError("Phiên Google hết hạn — đăng nhập lại")
        if resp.status_code in (403, 404):
            raise SheetAccessError(f"Tài khoản của bạn không mở được {what} — kiểm tra link và nhờ chủ file "
                                   "share quyền Người xem cho email của bạn")

    def _get(self, url: str, params: dict[str, str]) -> dict[str, Any]:
        resp = self._session.get(url, params=params, timeout=self._timeout)
        self._check(resp)
        if resp.status_code == 400:
            if _is_xlsx_error(resp):
                raise _NotANativeSheet()
            raise SheetAccessError("Không đọc được tab trong sheet — kiểm tra tên tab (mapping) hoặc gid trong link")
        if resp.status_code >= 400:
            raise SheetAccessError(f"Google Sheets API lỗi {resp.status_code}")
        data: dict[str, Any] = resp.json()
        return data

    # ------------------------------------------------------------------ file order (nhiều tab, .xlsx)
    def open(self, ref: SheetRef) -> OrderWorkbook:
        try:
            meta = self._get(f"{self.BASE}/{ref.spreadsheet_id}",
                             {"fields": "sheets.properties(sheetId,title,hidden)"})
        except _NotANativeSheet:
            return self._open_xlsx(ref.spreadsheet_id)
        except SheetAccessError:
            if ref.gid is not None:     # link docs.google.com của Google Sheet -> lỗi quyền thật
                raise
            try:                        # link drive.google.com/file/... có thể là file Excel
                return self._open_xlsx(ref.spreadsheet_id)
            except DriveScopeError:
                raise
            except SheetAccessError:
                pass
            raise
        props = [s["properties"] for s in meta.get("sheets", []) if not s["properties"].get("hidden")]
        if not props:
            raise SheetAccessError("Spreadsheet không có tab nào")
        gid_tab = next((p["title"] for p in props if p["sheetId"] == ref.gid), None) if ref.gid is not None else None

        def read(tab: str) -> list[list[Any]]:
            return self._values(ref.spreadsheet_id, tab)

        def heads(tabs: list[str]) -> dict[str, list[list[Any]]]:
            return self._batch(ref.spreadsheet_id, tabs, f"1:{HEAD_ROWS}")

        return OrderWorkbook("google_sheet", [p["title"] for p in props], read, heads, gid_tab)

    def _open_xlsx(self, file_id: str) -> OrderWorkbook:
        params = {"supportsAllDrives": "true"}
        resp = self._session.get(f"{self.DRIVE}/{file_id}", params={**params, "fields": "name,mimeType,size"},
                                 timeout=self._timeout)
        self._drive_check(resp)
        meta = resp.json()
        if meta.get("mimeType") != XLSX_MIME:
            raise SheetAccessError(f"File '{meta.get('name', '')}' không phải Google Sheet hay Excel (.xlsx)")
        if int(meta.get("size") or 0) > MAX_XLSX_BYTES:
            raise SheetAccessError("File Excel quá lớn (> 30 MB) — chuyển sang Google Sheet rồi dán link bản mới")
        resp = self._session.get(f"{self.DRIVE}/{file_id}", params={**params, "alt": "media"},
                                 timeout=max(self._timeout, 60.0))
        self._drive_check(resp)
        return xlsx_workbook(resp.content)

    def _drive_check(self, resp: Any) -> None:
        if resp.status_code == 401:
            raise SheetAuthError("Phiên Google hết hạn — đăng nhập lại")
        if resp.status_code == 403:
            body = resp.text if isinstance(getattr(resp, "text", None), str) else ""
            if "insufficient" in body.lower() or "scope" in body.lower():
                raise DriveScopeError("File order là Excel (.xlsx) trên Drive — cần cấp quyền đọc Drive: "
                                      "đăng xuất rồi đăng nhập lại và bấm Cho phép")
            if "accessNotConfigured" in body or "has not been used" in body or "disabled" in body:
                raise SheetAccessError("Google Drive API chưa bật trong project Google Cloud — báo IT bật "
                                       "(APIs & Services › Library › Google Drive API › Enable)")
        if resp.status_code in (403, 404):
            raise SheetAccessError("Tài khoản của bạn không mở được file — kiểm tra link và nhờ chủ file "
                                   "share quyền Người xem cho email của bạn")
        if resp.status_code >= 400:
            raise SheetAccessError(f"Google Drive API lỗi {resp.status_code}")

    def _values(self, spreadsheet_id: str, tab: str) -> list[list[Any]]:
        a1 = "'" + tab.replace("'", "''") + "'"
        data = self._get(f"{self.BASE}/{spreadsheet_id}/values/{quote(a1, safe='')}",
                         {"valueRenderOption": "UNFORMATTED_VALUE", "dateTimeRenderOption": "SERIAL_NUMBER",
                          "majorDimension": "ROWS"})
        values: list[list[Any]] = data.get("values", [])
        return values

    def _batch(self, spreadsheet_id: str, titles: list[str], rows: str = "") -> dict[str, list[list[Any]]]:
        if not titles:
            return {}
        suffix = f"!{rows}" if rows else ""
        params: list[tuple[str, str]] = [("ranges", "'" + t.replace("'", "''") + "'" + suffix) for t in titles]
        params += [("valueRenderOption", "UNFORMATTED_VALUE"), ("dateTimeRenderOption", "SERIAL_NUMBER"),
                   ("majorDimension", "ROWS")]
        resp = self._session.get(f"{self.BASE}/{spreadsheet_id}/values:batchGet", params=params,
                                 timeout=self._timeout)
        self._check(resp)
        if resp.status_code >= 400:
            raise SheetAccessError(f"Google Sheets API lỗi {resp.status_code}")
        ranges = resp.json().get("valueRanges", [])
        return {t: r.get("values", []) for t, r in zip(titles, ranges, strict=False)}

    def read(self, ref: SheetRef, *, sheet_name: str | None = None) -> list[list[Any]]:
        try:
            return self._read_one(ref, sheet_name)
        except _NotANativeSheet as exc:
            raise SheetAccessError(XLSX_HINT) from exc

    def _read_one(self, ref: SheetRef, sheet_name: str | None) -> list[list[Any]]:
        title = sheet_name
        if ref.gid is not None or title is None:
            meta = self._get(f"{self.BASE}/{ref.spreadsheet_id}", {"fields": "sheets.properties(sheetId,title)"})
            sheets = [s["properties"] for s in meta.get("sheets", [])]
            if not sheets:
                raise SheetAccessError("Spreadsheet không có tab nào")
            if ref.gid is not None:
                match = next((s for s in sheets if s["sheetId"] == ref.gid), None)
                if match is None:
                    raise SheetAccessError(f"Không có tab gid={ref.gid}")
                title = match["title"]
            else:
                title = sheets[0]["title"]
        return self._values(ref.spreadsheet_id, title)

    def titles(self, ref: SheetRef) -> list[str]:
        meta = self._get(f"{self.BASE}/{ref.spreadsheet_id}", {"fields": "sheets.properties(title)"})
        return [s["properties"]["title"] for s in meta.get("sheets", [])]

    def read_tabs(self, ref: SheetRef, titles: list[str]) -> dict[str, list[list[Any]]]:
        """Đọc nhiều tab trong 1 request (values:batchGet) — dùng cho Sheet master."""
        try:
            return self._batch(ref.spreadsheet_id, titles)
        except SheetAccessError as exc:
            if isinstance(exc, SheetAuthError) or "không mở được" not in str(exc):
                raise
            raise SheetAccessError(str(exc).replace("không mở được sheet", "không mở được Sheet master")) from exc


XLSX_HINT = ("Link này là file Excel (.xlsx) trên Drive, không phải Google Sheet — mở file › Tệp › Lưu dưới dạng "
             "Google Trang tính, rồi dán link của bản mới")


class _NotANativeSheet(SheetAccessError):
    """Sheets API trả 400 'not supported for this document' -> file Excel/khác trên Drive."""

    def __init__(self) -> None:
        super().__init__(XLSX_HINT)


def _is_xlsx_error(resp: Any) -> bool:
    body = resp.text if isinstance(getattr(resp, "text", None), str) else ""
    return "not supported for this document" in body


def read_xlsx_tabs(path: str, titles: list[str] | None = None) -> dict[str, list[list[Any]]]:
    """Đọc các tab của file Excel (dùng cho CLI / test: Sheet master tải về dạng .xlsx)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = wb.sheetnames if titles is None else [t for t in titles if t in wb.sheetnames]
        return {n: _trim_rows([list(r) for r in wb[n].iter_rows(values_only=True)]) for n in names}
    finally:
        wb.close()


def xlsx_sheet_names(path: str) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
