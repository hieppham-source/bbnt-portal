"""
BBNT Calculation Engine — parse file đối soát, tính NMV, hoa hồng và thuế.

Pipeline::

    rows (Excel / Google Sheets values)
      └─ parse_sheet(rows, MappingConfig)        -> ParseResult(lines, issues)
           └─ compute_nmv(lines, period_months)  -> NmvResult(A = GMV ước tính, B = NMV "Đã giao")
                └─ calculate_commission(B, rate, TaxPolicy) -> CommissionResult

    reconcile(...) chạy cả 3 bước.

Quy tắc nghiệp vụ:

* **NMV (B)** chỉ cộng dòng có status chuẩn hoá = ``Đã giao``. ``Hoàn hàng``, ``Hủy đơn`` và
  status lạ/trống bị loại khỏi NMV và sinh cảnh báo.
* **GMV ước tính (A)** cộng mọi dòng hợp lệ trong kỳ (mọi status) -> tỷ lệ huỷ = 1 - B/A.
* **EZS**:   Revenue = Base x Rate; Total = Revenue + VAT 8%.
* **SASSY**: Revenue = Base x Rate; Total = Revenue - 1% khấu trừ thuế.
* Mọi số tiền là ``Decimal`` (không float), làm tròn ROUND_HALF_UP tới ``quantum`` (mặc định 0.01).
"""
from __future__ import annotations

import enum
import re
import unicodedata
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from .numbering import VN_TZ, YearMonth, YearMonthLike, normalize_months

Number = Decimal | int | str
ZERO = Decimal("0")
CENT = Decimal("0.01")


# =============================================================================
# Errors & issues
# =============================================================================
class MappingError(ValueError):
    """Cấu hình mapping không khớp file (thiếu cột bắt buộc, header_row sai...)."""


class TaxRuleNotFoundError(LookupError):
    """Không có quy tắc thuế hiệu lực cho pháp nhân tại ngày yêu cầu."""


class Severity(str, enum.Enum):
    WARNING = "WARNING"  # dòng bị loại khỏi NMV theo đúng nghiệp vụ — cần người xem
    ERROR = "ERROR"      # dữ liệu hỏng — chặn phát hành BBNT tới khi sửa


class IssueCode(str, enum.Enum):
    RETURNED = "RETURNED"
    CANCELLED = "CANCELLED"
    UNKNOWN_STATUS = "UNKNOWN_STATUS"
    OUT_OF_PERIOD = "OUT_OF_PERIOD"
    DUPLICATE_ROW = "DUPLICATE_ROW"
    MISSING_REQUIRED = "MISSING_REQUIRED"
    INVALID_VALUE = "INVALID_VALUE"
    NEGATIVE_AMOUNT = "NEGATIVE_AMOUNT"


_SEVERITY: dict[IssueCode, Severity] = {
    IssueCode.RETURNED: Severity.WARNING,
    IssueCode.CANCELLED: Severity.WARNING,
    IssueCode.UNKNOWN_STATUS: Severity.WARNING,
    IssueCode.OUT_OF_PERIOD: Severity.WARNING,
    IssueCode.DUPLICATE_ROW: Severity.WARNING,
    IssueCode.MISSING_REQUIRED: Severity.ERROR,
    IssueCode.INVALID_VALUE: Severity.ERROR,
    IssueCode.NEGATIVE_AMOUNT: Severity.ERROR,
}


@dataclass(frozen=True)
class RowIssue:
    """Một vấn đề gắn với 1 dòng file (``row_number`` 1-based như Excel/Sheets hiển thị)."""

    row_number: int
    code: IssueCode
    message: str
    order_code: str | None = None
    amount: Decimal | None = None

    @property
    def severity(self) -> Severity:
        return _SEVERITY[self.code]


# =============================================================================
# Text normalisation (tiếng Việt)
# =============================================================================
def normalize_text(value: Any) -> str:
    """Chuẩn hoá để so khớp: bỏ dấu, đ->d, casefold, gộp khoảng trắng.

    Xử lý được: ``'Hủy'`` vs ``'Huỷ'`` (vị trí dấu khác nhau), chuỗi NFC/NFD, NBSP, hoa/thường.
    """
    if value is None:
        return ""
    s = unicodedata.normalize("NFD", str(value))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").casefold()
    return " ".join(s.split())


# =============================================================================
# Status
# =============================================================================
class OrderStatus(str, enum.Enum):
    DELIVERED = "DELIVERED"
    RETURNED = "RETURNED"
    CANCELLED = "CANCELLED"
    UNKNOWN = "UNKNOWN"


#: Alias mặc định. Có thể override per-mapping qua ``reconciliation_mappings.options.status_aliases``.
DEFAULT_STATUS_ALIASES: Mapping[OrderStatus, frozenset[str]] = {
    OrderStatus.DELIVERED: frozenset({"Đã giao", "Đã giao hàng", "Giao thành công", "Giao hàng thành công"}),
    OrderStatus.RETURNED: frozenset({"Hoàn hàng", "Đã hoàn hàng", "Trả hàng", "Hoàn trả", "Đã hoàn"}),
    OrderStatus.CANCELLED: frozenset({"Hủy đơn", "Đã hủy", "Hủy", "Hủy đơn hàng", "Yêu cầu hủy", "Khách hủy"}),
}


class StatusClassifier:
    """Map status thô -> :class:`OrderStatus` (so khớp sau :func:`normalize_text`)."""

    def __init__(self, aliases: Mapping[OrderStatus, Iterable[str]] = DEFAULT_STATUS_ALIASES) -> None:
        self._lookup: dict[str, OrderStatus] = {}
        for status, names in aliases.items():
            if status is OrderStatus.UNKNOWN:
                raise ValueError("Không khai báo alias cho UNKNOWN")
            for name in names:
                key = normalize_text(name)
                if not key:
                    continue
                if self._lookup.get(key, status) is not status:
                    raise ValueError(f"Alias '{name}' bị gán cho 2 status")
                self._lookup[key] = status

    def classify(self, raw: Any) -> OrderStatus:
        return self._lookup.get(normalize_text(raw), OrderStatus.UNKNOWN)


# =============================================================================
# Value parsers
# =============================================================================
_CURRENCY_RE = re.compile(r"(?i)(vnd|vnđ|đ|₫|\s| )")


def parse_money(value: Any) -> Decimal | None:
    """Parse số tiền từ Excel/Sheets -> ``Decimal`` (``None`` nếu ô trống).

    * số Excel (int/float) -> Decimal qua ``str()`` để không mang sai số nhị phân
    * ``'1.250.000'``, ``'1,250,000'``, ``'1.250.000 ₫'``, ``'463500'`` -> 1250000 / 463500
    * ``'1.234,56'`` (VN), ``'1,234.56'`` (US) -> 1234.56
    * ``'1.250'`` -> 1250 (1 dấu + đúng 3 chữ số => phân cách nghìn — VND không có lẻ 3 số)
    * ``'(1.000)'``, ``'-1.000'`` -> -1000

    Raises:
        ValueError: chuỗi không phải số / nhóm nghìn sai.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f"Giá trị boolean không phải số tiền: {value!r}")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            raise ValueError(f"Số không hợp lệ: {value!r}")
        return Decimal(str(value))
    s = _CURRENCY_RE.sub("", str(value))
    if not s or s in {"-", "—"}:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    if s.startswith("-"):
        negative, s = not negative, s[1:]
    if not re.fullmatch(r"[0-9.,]+", s) or not re.search(r"\d", s):
        raise ValueError(f"Không đọc được số tiền: {value!r}")

    if "." in s and "," in s:
        dec = "." if s.rfind(".") > s.rfind(",") else ","
        thousands = "," if dec == "." else "."
        int_part, frac = s.rsplit(dec, 1)
        _check_groups(int_part.split(thousands), value)
        s = int_part.replace(thousands, "") + "." + frac
    elif "." in s or "," in s:
        sep = "." if "." in s else ","
        parts = s.split(sep)
        if len(parts) > 2 or len(parts[1]) == 3:
            _check_groups(parts, value)
            s = "".join(parts)
        else:
            s = parts[0] + "." + parts[1]
    try:
        result = Decimal(s)
    except InvalidOperation as exc:  # pragma: no cover - regex đã chặn
        raise ValueError(f"Không đọc được số tiền: {value!r}") from exc
    return -result if negative else result


def _check_groups(groups: Sequence[str], original: Any) -> None:
    if not groups[0] or len(groups[0]) > 3 or any(len(g) != 3 for g in groups[1:]):
        raise ValueError(f"Phân cách hàng nghìn không hợp lệ: {original!r}")


def parse_quantity(value: Any) -> int | None:
    q = parse_money(value)
    if q is None:
        return None
    if q != q.to_integral_value():
        raise ValueError(f"Số lượng phải là số nguyên: {value!r}")
    return int(q)


_EXCEL_EPOCH = datetime(1899, 12, 30)
_DT_FORMATS = ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
               "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
               "%Y/%m/%d %H:%M:%S", "%Y/%m/%d")


_MERIDIEM_RE = re.compile(r"(?i)\s*(SA|CH|AM|PM)\.?$")


def _split_meridiem(s: str) -> tuple[str, str | None]:
    """Tách hậu tố buổi: 'SA'/'AM' (sáng) | 'CH'/'PM' (chiều) — Excel locale VN xuất '21:53 CH'."""
    m = _MERIDIEM_RE.search(s)
    if not m:
        return s, None
    return s[:m.start()], "AM" if m.group(1).upper() in ("SA", "AM") else "PM"


def _parse_datetime_string(raw: str) -> datetime:
    s, meridiem = _split_meridiem(raw.strip())
    parsed: datetime | None = None
    try:
        parsed = datetime.fromisoformat(s)
    except ValueError:
        for fmt in _DT_FORMATS:
            try:
                parsed = datetime.strptime(s, fmt)
                break
            except ValueError:
                continue
    if parsed is None:
        raise ValueError(f"Không đọc được ngày đặt hàng: {raw!r}")
    if meridiem == "PM" and parsed.hour < 12:
        parsed += timedelta(hours=12)
    elif meridiem == "AM" and parsed.hour == 12:
        parsed -= timedelta(hours=12)
    return parsed


def parse_order_datetime(value: Any, naive_tz: ZoneInfo = VN_TZ) -> datetime | None:
    """Parse thời điểm đặt hàng -> datetime aware theo **giờ VN**.

    * datetime aware (vd ISO ``...Z`` từ API TikTok) -> đổi sang VN:
      ``2026-05-31T17:30:00Z`` thành ``2026-06-01 00:30 +07`` (thuộc tháng 6!)
    * datetime/chuỗi naive -> hiểu theo ``naive_tz`` (mặc định VN; set UTC cho file export UTC)
    * số (serial date của Excel / Google Sheets UNFORMATTED_VALUE) -> từ mốc 1899-12-30
    * chuỗi ``dd/mm/yyyy [HH:MM[:SS]]`` (kiểu VN — ngày đứng trước) hoặc ISO 8601
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime.combine(value, time())
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        dt = _EXCEL_EPOCH + timedelta(seconds=round(float(value) * 86400))  # làm tròn tới giây
    else:
        dt = _parse_datetime_string(str(value))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=naive_tz)
    return dt.astimezone(VN_TZ)


# =============================================================================
# Mapping
# =============================================================================
class TargetField(str, enum.Enum):
    """Đồng bộ enum ``recon_target_field`` trong DB."""

    ORDER_DATE = "ORDER_DATE"
    ORDER_CODE = "ORDER_CODE"
    ORDER_STATUS = "ORDER_STATUS"
    BRAND_NAME = "BRAND_NAME"
    SKU = "SKU"
    PRODUCT_NAME = "PRODUCT_NAME"
    QUANTITY = "QUANTITY"
    UNIT_PRICE = "UNIT_PRICE"
    DISCOUNT = "DISCOUNT"
    GMV = "GMV"
    GMV_AFTER_VOUCHER = "GMV_AFTER_VOUCHER"
    CAMPAIGN_NAME = "CAMPAIGN_NAME"
    PAYMENT_METHOD = "PAYMENT_METHOD"
    NOTE = "NOTE"


@dataclass(frozen=True)
class FieldSpec:
    """Ánh xạ 1 field: ưu tiên ``header`` (so khớp không dấu/không hoa thường), fallback ``column`` (A, B, AA)."""

    target: TargetField
    header: str | None = None
    column: str | None = None
    required: bool = False

    def __post_init__(self) -> None:
        if not self.header and not self.column:
            raise MappingError(f"{self.target.value}: cần header hoặc column")
        if self.column and not re.fullmatch(r"[A-Z]{1,3}", self.column):
            raise MappingError(f"{self.target.value}: column '{self.column}' không hợp lệ")


@dataclass(frozen=True)
class MappingConfig:
    """Bản runtime của ``reconciliation_mappings`` + ``reconciliation_mapping_fields``."""

    fields: tuple[FieldSpec, ...]
    header_row: int = 1
    data_start_row: int = 2
    sheet_name: str | None = None
    status_aliases: Mapping[OrderStatus, frozenset[str]] = field(default_factory=lambda: DEFAULT_STATUS_ALIASES)
    naive_tz: ZoneInfo = VN_TZ

    def __post_init__(self) -> None:
        if self.header_row < 1 or self.data_start_row <= self.header_row:
            raise MappingError("Cần 1 <= header_row < data_start_row")
        targets = [f.target for f in self.fields]
        dupes = [t.value for t, n in Counter(targets).items() if n > 1]
        if dupes:
            raise MappingError(f"Field bị map 2 lần: {', '.join(dupes)}")
        present = set(targets)
        missing = {TargetField.ORDER_CODE, TargetField.ORDER_STATUS} - present
        if missing:
            raise MappingError("Mapping bắt buộc có: " + ", ".join(sorted(m.value for m in missing))
                               + " (NMV chỉ tính đơn 'Đã giao' nên phải có cột status)")
        if not ({TargetField.GMV_AFTER_VOUCHER, TargetField.GMV} & present
                or {TargetField.QUANTITY, TargetField.UNIT_PRICE} <= present):
            raise MappingError("Mapping cần GMV_AFTER_VOUCHER, GMV, hoặc cặp QUANTITY + UNIT_PRICE")

    @classmethod
    def from_db(cls, mapping: Any) -> MappingConfig:
        """Dựng từ ORM ``ReconciliationMapping`` (có ``.fields``, ``.options``)."""
        opts: Mapping[str, Any] = mapping.options or {}
        aliases = DEFAULT_STATUS_ALIASES
        if "status_aliases" in opts:
            aliases = {OrderStatus(k): frozenset(v) for k, v in opts["status_aliases"].items()}
        return cls(
            fields=tuple(FieldSpec(TargetField(f.target_field.value if hasattr(f.target_field, "value")
                                               else f.target_field),
                                   header=f.source_header, column=f.source_column, required=f.is_required)
                         for f in mapping.fields),
            header_row=mapping.header_row,
            data_start_row=mapping.data_start_row,
            sheet_name=mapping.sheet_name,
            status_aliases=aliases,
            naive_tz=ZoneInfo(opts.get("naive_timezone", "Asia/Ho_Chi_Minh")),
        )


def column_index(letter: str) -> int:
    """'A' -> 0, 'Z' -> 25, 'AA' -> 26."""
    idx = 0
    for ch in letter:
        idx = idx * 26 + (ord(ch) - 64)
    return idx - 1


#: Tên cột hay gặp trong file order của các brand — dùng SAU tên cột khai trong MAPPING (chỉ cho field có trong
#: mapping). Brand đặt tên lệch chuẩn vẫn đọc được mà không phải sửa Sheet master mỗi lần.
DEFAULT_HEADER_ALIASES: dict[TargetField, tuple[str, ...]] = {
    TargetField.ORDER_DATE: ("Ngày đặt hàng", "Ngày đặt", "Ngày tạo đơn", "Ngày đơn hàng", "Order date",
                             "Created time"),
    TargetField.ORDER_CODE: ("Mã đơn hàng", "Mã đơn", "Order ID", "Order code", "Order number"),
    TargetField.ORDER_STATUS: ("Trạng thái", "Trạng thái đơn hàng", "Trạng thái đơn", "Tình trạng lên đơn",
                               "Tình trạng đơn hàng", "Tình trạng đơn", "Tình trạng", "ORDER_STATUS",
                               "Order status", "Status"),
    TargetField.BRAND_NAME: ("Hãng", "Brand", "Thương hiệu", "Nhãn hàng"),
    TargetField.SKU: ("Mã sản phẩm", "SKU", "Mã SP", "Mã hàng"),
    TargetField.PRODUCT_NAME: ("Tên sản phẩm", "Sản phẩm", "Tên hàng", "Product name"),
    TargetField.QUANTITY: ("Số lượng", "SL", "Quantity", "Qty"),
    TargetField.UNIT_PRICE: ("Giá sản phẩm", "Đơn giá", "Giá bán", "Unit price"),
    TargetField.DISCOUNT: ("Số tiền giảm", "Giảm giá", "Discount"),
    TargetField.GMV: ("Total GMV", "GMV", "Tổng tiền", "Thành tiền"),
    TargetField.GMV_AFTER_VOUCHER: ("GMV sau voucher (if any)", "GMV sau voucher", "Doanh thu sau voucher",
                                    "Doanh thu sau giảm giá", "GMV after voucher"),
    TargetField.CAMPAIGN_NAME: ("Tên Camp", "Tên campaign", "Campaign", "Chiến dịch"),
    TargetField.PAYMENT_METHOD: ("Phương thức thanh toán", "Payment method"),
    TargetField.NOTE: ("Ghi chú", "Note"),
}

_ALIAS_SPLIT = re.compile(r"\s*(?:\||\n)\s*")
_HEADER_NOISE = re.compile(r"[*:]+")


def header_aliases(header: str | None) -> list[str]:
    """'Trạng thái | Tình trạng lên đơn' -> ['Trạng thái', 'Tình trạng lên đơn'] (ngăn bằng | hoặc xuống dòng)."""
    return [a for a in _ALIAS_SPLIT.split(header or "") if a.strip()]


def _norm_header(value: Any) -> str:
    return " ".join(_HEADER_NOISE.sub(" ", normalize_text(value)).split())


def _find_header(names: Sequence[str], header_norm: Sequence[str]) -> int | None:
    """Tìm cột theo danh sách tên (ưu tiên theo thứ tự): khớp nguyên văn trước, sau đó 'tên cột chứa alias'
    nếu chỉ đúng 1 cột chứa (vd 'BRAND INPUT Tình trạng lên đơn' chứa 'Tình trạng lên đơn')."""
    wanted = [n for n in (_norm_header(x) for x in names) if n]
    for w in wanted:
        if w in header_norm:
            return header_norm.index(w)
    for w in wanted:
        if len(w) < 4:
            continue
        hits = [i for i, h in enumerate(header_norm) if h and re.search(rf"(?:^|\W){re.escape(w)}(?:\W|$)", h)]
        if len(hits) == 1:
            return hits[0]
    return None


def resolve_columns(header: Sequence[Any], config: MappingConfig, *, header_row: int | None = None,
                    ) -> dict[TargetField, int]:
    """Tìm vị trí cột cho từng field. Header trùng tên -> lấy cột đầu tiên.

    Thứ tự thử: các tên khai ở MAPPING (ngăn bằng ``|``) -> tên phổ biến :data:`DEFAULT_HEADER_ALIASES`
    -> cột chữ cái (A, B, AA) nếu MAPPING có khai.

    Raises:
        MappingError: field ``required``/bắt buộc không tìm thấy.
    """
    header_norm = [_norm_header(c) for c in header]
    hard_required = {TargetField.ORDER_CODE, TargetField.ORDER_STATUS}
    resolved: dict[TargetField, int] = {}
    required_missing: list[str] = []
    for spec in config.fields:
        idx = _find_header(header_aliases(spec.header), header_norm) if spec.header else None
        if idx is None:
            idx = _find_header(DEFAULT_HEADER_ALIASES.get(spec.target, ()), header_norm)
        if idx is None and spec.column:
            idx = column_index(spec.column)
        if idx is not None and idx not in resolved.values():
            resolved[spec.target] = idx
        elif spec.required or spec.target in hard_required:
            required_missing.append(f"{spec.target.value} ('{spec.header or spec.column}')")

    amount_ok = ({TargetField.GMV_AFTER_VOUCHER, TargetField.GMV} & resolved.keys()
                 or {TargetField.QUANTITY, TargetField.UNIT_PRICE} <= resolved.keys())
    if required_missing or not amount_ok:
        detail = ", ".join(required_missing) or "cột số tiền (GMV / GMV sau voucher / SL + đơn giá)"
        raise MappingError(f"Không tìm thấy cột trong header dòng {header_row or config.header_row}: {detail}")
    return resolved


#: Số dòng đầu file được dò tìm header khi dòng header khai trong MAPPING không khớp.
HEADER_SCAN_ROWS = 15


def locate_header(rows: Sequence[Sequence[Any]], config: MappingConfig) -> tuple[int, dict[TargetField, int]]:
    """Trả về ``(dòng header 1-based, vị trí cột)``. Thử dòng khai trong MAPPING trước, sau đó dò
    ``HEADER_SCAN_ROWS`` dòng đầu (file brand hay chèn thêm dòng tổng / tiêu đề phía trên header)."""
    if len(rows) < config.header_row:
        raise MappingError(f"File chỉ có {len(rows)} dòng, không tới header_row={config.header_row}")
    try:
        return config.header_row, resolve_columns(rows[config.header_row - 1], config)
    except MappingError as first:
        for i in range(min(len(rows), HEADER_SCAN_ROWS)):
            if i == config.header_row - 1:
                continue
            try:
                return i + 1, resolve_columns(rows[i], config, header_row=i + 1)
            except MappingError:
                continue
        raise first


# =============================================================================
# Parsing
# =============================================================================
@dataclass(frozen=True)
class OrderLine:
    row_number: int
    order_code: str
    status: OrderStatus
    raw_status: str
    amount: Decimal
    quantity: int | None = None
    sku: str | None = None
    product_name: str | None = None
    brand_name: str | None = None
    order_date: datetime | None = None  # giờ VN


@dataclass(frozen=True)
class ParseResult:
    lines: tuple[OrderLine, ...]
    issues: tuple[RowIssue, ...]
    rows_read: int
    header_row: int = 1


def _cell(row: Sequence[Any], idx: int | None) -> Any:
    """Lấy ô an toàn — Google Sheets API cắt bỏ ô trống cuối dòng nên row có thể ngắn hơn header."""
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    return v.strip() if isinstance(v, str) else v


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def parse_sheet(rows: Sequence[Sequence[Any]], config: MappingConfig) -> ParseResult:
    """Parse ma trận ô (``list[list]``) thành :class:`OrderLine`.

    ``rows`` có thể lấy từ :func:`rows_from_xlsx` hoặc trực tiếp ``values`` của Google Sheets API.
    Dòng trống hoàn toàn bị bỏ qua. Dòng lỗi được ghi :class:`RowIssue` (ERROR) và không vào kết quả.

    Số tiền 1 dòng = ``GMV_AFTER_VOUCHER`` nếu có giá trị, ngược lại ``GMV``, ngược lại
    ``QUANTITY x UNIT_PRICE - DISCOUNT``.
    """
    header_row, cols = locate_header(rows, config)
    data_start_row = header_row + (config.data_start_row - config.header_row)
    classifier = StatusClassifier(config.status_aliases)
    F = TargetField
    identity = [cols[t] for t in (F.ORDER_CODE, F.ORDER_DATE, F.SKU, F.PRODUCT_NAME) if t in cols]

    lines: list[OrderLine] = []
    issues: list[RowIssue] = []
    rows_read = 0
    for offset, row in enumerate(rows[data_start_row - 1:]):
        rn = data_start_row + offset
        if all(_blank(c) for c in row):
            continue
        if all(_blank(_cell(row, i)) for i in identity):
            continue    # dòng kẻ sẵn công thức (vd GMV = 0) nhưng chưa có đơn -> coi như dòng trống
        rows_read += 1
        order_code = _cell(row, cols.get(F.ORDER_CODE))
        order_code = str(order_code) if not _blank(order_code) else None
        if order_code is None:
            issues.append(RowIssue(rn, IssueCode.MISSING_REQUIRED, "Thiếu mã đơn hàng"))
            continue
        try:
            amount = _line_amount(row, cols)
            quantity = parse_quantity(_cell(row, cols.get(F.QUANTITY)))
            order_date = parse_order_datetime(_cell(row, cols.get(F.ORDER_DATE)), config.naive_tz)
        except ValueError as exc:
            issues.append(RowIssue(rn, IssueCode.INVALID_VALUE, str(exc), order_code))
            continue
        if amount is None:
            issues.append(RowIssue(rn, IssueCode.MISSING_REQUIRED, "Thiếu số tiền", order_code))
            continue
        if amount < 0:
            issues.append(RowIssue(rn, IssueCode.NEGATIVE_AMOUNT, f"Số tiền âm: {amount}", order_code, amount))
            continue
        raw_status = _cell(row, cols.get(F.ORDER_STATUS))
        lines.append(OrderLine(
            row_number=rn,
            order_code=order_code,
            status=classifier.classify(raw_status),
            raw_status="" if raw_status is None else str(raw_status),
            amount=amount,
            quantity=quantity,
            sku=_str_or_none(_cell(row, cols.get(F.SKU))),
            product_name=_str_or_none(_cell(row, cols.get(F.PRODUCT_NAME))),
            brand_name=_str_or_none(_cell(row, cols.get(F.BRAND_NAME))),
            order_date=order_date,
        ))
    return ParseResult(tuple(lines), tuple(issues), rows_read, header_row)


def _str_or_none(v: Any) -> str | None:
    return None if _blank(v) else str(v)


def _line_amount(row: Sequence[Any], cols: Mapping[TargetField, int]) -> Decimal | None:
    F = TargetField
    for target in (F.GMV_AFTER_VOUCHER, F.GMV):
        v = parse_money(_cell(row, cols.get(target)))
        if v is not None:
            return v
    qty = parse_money(_cell(row, cols.get(F.QUANTITY)))
    price = parse_money(_cell(row, cols.get(F.UNIT_PRICE)))
    if qty is None or price is None:
        return None
    return qty * price - (parse_money(_cell(row, cols.get(F.DISCOUNT))) or ZERO)


def rows_from_xlsx(path: str | Path, sheet_name: str | None = None) -> list[list[Any]]:
    """Đọc sheet Excel thành ``list[list]`` (giá trị đã tính, không phải công thức)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


# =============================================================================
# NMV
# =============================================================================
@dataclass(frozen=True)
class BbntItemRow:
    """1 dòng bảng TÊN SẢN PHẨM | SỐ LƯỢNG | DOANH THU (chỉ đơn Đã giao)."""

    brand_name: str | None
    sku: str | None
    product_name: str
    quantity: int
    revenue: Decimal


@dataclass(frozen=True)
class NmvResult:
    gmv_estimated: Decimal                       # (A) mọi dòng hợp lệ trong kỳ
    nmv: Decimal                                 # (B) chỉ "Đã giao"
    cancellation_rate: Decimal | None         # 1 - B/A, 6 chữ số; None nếu A = 0
    totals_by_status: Mapping[OrderStatus, Decimal]
    counts_by_status: Mapping[OrderStatus, int]
    items: tuple[BbntItemRow, ...]
    issues: tuple[RowIssue, ...]


_STATUS_ISSUE = {
    OrderStatus.RETURNED: (IssueCode.RETURNED, "Hoàn hàng — không tính NMV"),
    OrderStatus.CANCELLED: (IssueCode.CANCELLED, "Hủy đơn — không tính NMV"),
    OrderStatus.UNKNOWN: (IssueCode.UNKNOWN_STATUS, "Status chưa rõ — không tính NMV, cần kiểm tra"),
}


def compute_nmv(lines: Iterable[OrderLine],
                period_months: Iterable[YearMonthLike] | None = None) -> NmvResult:
    """Tính A (GMV ước tính), B (NMV) và bảng sản phẩm.

    Args:
        lines: Kết quả :func:`parse_sheet`.
        period_months: Nếu truyền, chỉ giữ dòng có ``order_date`` (giờ VN) thuộc các tháng này.
            Dòng ngoài kỳ -> ``OUT_OF_PERIOD``; dòng thiếu ngày -> ``MISSING_REQUIRED`` (ERROR).
    """
    months = set(normalize_months(period_months)) if period_months is not None else None
    totals: dict[OrderStatus, Decimal] = {s: ZERO for s in OrderStatus}
    counts: dict[OrderStatus, int] = {s: 0 for s in OrderStatus}
    issues: list[RowIssue] = []
    items: dict[tuple[str | None, str | None, str], list[Any]] = {}
    seen: Counter[tuple[Any, ...]] = Counter()

    for ln in lines:
        if months is not None:
            if ln.order_date is None:
                issues.append(RowIssue(ln.row_number, IssueCode.MISSING_REQUIRED,
                                       "Thiếu ngày đặt hàng, không xác định được kỳ", ln.order_code, ln.amount))
                continue
            if YearMonth.of(ln.order_date) not in months:
                issues.append(RowIssue(ln.row_number, IssueCode.OUT_OF_PERIOD,
                                       f"Đơn ngày {ln.order_date:%d/%m/%Y %H:%M} (giờ VN) ngoài kỳ",
                                       ln.order_code, ln.amount))
                continue

        fingerprint = (ln.order_code, ln.sku, ln.product_name, ln.quantity, ln.amount, ln.raw_status)
        seen[fingerprint] += 1
        if seen[fingerprint] == 2:
            issues.append(RowIssue(ln.row_number, IssueCode.DUPLICATE_ROW,
                                   "Dòng trùng hoàn toàn với dòng trước (vẫn được tính) — kiểm tra file gộp",
                                   ln.order_code, ln.amount))

        totals[ln.status] += ln.amount
        counts[ln.status] += 1
        if ln.status is OrderStatus.DELIVERED:
            key = (ln.brand_name, ln.sku, ln.product_name or ln.sku or "")
            agg = items.setdefault(key, [0, ZERO])
            agg[0] += ln.quantity or 0
            agg[1] += ln.amount
        else:
            code, msg = _STATUS_ISSUE[ln.status]
            detail = f"{msg} (status: '{ln.raw_status}')" if ln.raw_status else msg
            issues.append(RowIssue(ln.row_number, code, detail, ln.order_code, ln.amount))

    gmv = sum(totals.values(), ZERO)
    nmv = totals[OrderStatus.DELIVERED]
    rate = (Decimal(1) - nmv / gmv).quantize(Decimal("0.000001"), ROUND_HALF_UP) if gmv else None
    return NmvResult(
        gmv_estimated=gmv,
        nmv=nmv,
        cancellation_rate=rate,
        totals_by_status=totals,
        counts_by_status=counts,
        items=tuple(BbntItemRow(k[0], k[1], k[2], v[0], v[1]) for k, v in items.items()),
        issues=tuple(sorted(issues, key=lambda i: i.row_number)),
    )


# =============================================================================
# Tax & commission
# =============================================================================
class TaxMode(str, enum.Enum):
    VAT_ADD = "VAT_ADD"        # cộng VAT (công ty)
    TAX_DEDUCT = "TAX_DEDUCT"  # trừ thuế khấu trừ (hộ kinh doanh)
    NONE = "NONE"


@dataclass(frozen=True)
class TaxPolicy:
    mode: TaxMode
    rate: Decimal
    label: str
    effective_from: date | None = None
    effective_to: date | None = None  # inclusive

    def __post_init__(self) -> None:
        if not isinstance(self.rate, Decimal):
            raise TypeError("TaxPolicy.rate phải là Decimal")
        if not ZERO <= self.rate < 1:
            raise ValueError(f"Thuế suất ngoài [0, 1): {self.rate}")
        if self.mode is TaxMode.NONE and self.rate != 0:
            raise ValueError("TaxMode.NONE phải có rate = 0")

    def applies_on(self, d: date) -> bool:
        return ((self.effective_from is None or self.effective_from <= d)
                and (self.effective_to is None or d <= self.effective_to))

    @classmethod
    def from_db(cls, rule: Any) -> TaxPolicy:
        """Dựng từ ORM ``EntityTaxRule``."""
        mode = rule.tax_mode.value if hasattr(rule.tax_mode, "value") else rule.tax_mode
        return cls(TaxMode(mode), Decimal(rule.tax_rate), rule.label, rule.effective_from, rule.effective_to)


#: Mặc định khi chưa đọc từ bảng ``entity_tax_rules``. NQ 204/2025/QH15 hiệu lực tới 31/12/2026 —
#: sau đó KHÔNG đoán thuế suất mà raise, buộc cấu hình lại.
DEFAULT_TAX_RULES: Mapping[str, tuple[TaxPolicy, ...]] = {
    "EZS": (TaxPolicy(TaxMode.VAT_ADD, Decimal("0.08"), "VAT 8%", date(2025, 7, 1), date(2026, 12, 31)),),
    "SASSY": (TaxPolicy(TaxMode.TAX_DEDUCT, Decimal("0.01"),
                        "Số tiền giảm tương ứng 20% mức tỷ lệ % để tính thuế GTGT theo NQ 204/2025/QH15",
                        date(2025, 7, 1), date(2026, 12, 31)),),
}
_ENTITY_ALIASES = {"EZ": "EZS", "SS": "SASSY"}


def resolve_tax_policy(entity: str, on: date,
                       rules: Mapping[str, Sequence[TaxPolicy]] = DEFAULT_TAX_RULES) -> TaxPolicy:
    """Chọn quy tắc thuế hiệu lực tại ngày ``on`` (nên dùng ngày LẬP BBNT theo giờ VN)."""
    key = entity.strip().upper()
    key = _ENTITY_ALIASES.get(key, key)
    for policy in rules.get(key, ()):
        if policy.applies_on(on):
            return policy
    raise TaxRuleNotFoundError(f"Chưa cấu hình thuế cho {entity} tại ngày {on:%d/%m/%Y}")


@dataclass(frozen=True)
class CommissionResult:
    """Khớp các cột ``generated_bbnts``; ``total_payable`` = A + B - C - D + F."""

    commission_base: Decimal
    commission_rate: Decimal
    revenue: Decimal                 # A — hoa hồng chưa thuế
    tax_mode: TaxMode
    tax_rate: Decimal
    vat_amount: Decimal              # B — EZS
    tax_deduction_amount: Decimal    # C — SASSY
    collected_on_behalf: Decimal     # D
    consulting_fee: Decimal          # F
    total_payable: Decimal           # E


def _to_decimal(value: Number, name: str) -> Decimal:
    if isinstance(value, float):
        raise TypeError(f"{name}: không nhận float (sai số nhị phân) — dùng Decimal hoặc str")
    if isinstance(value, bool):
        raise TypeError(f"{name}: không nhận bool")
    try:
        return Decimal(value)
    except (InvalidOperation, TypeError) as exc:
        raise ValueError(f"{name}: không phải số: {value!r}") from exc


def calculate_commission(
    commission_base: Number,
    commission_rate: Number,
    tax: TaxPolicy,
    *,
    collected_on_behalf: Number = 0,
    consulting_fee: Number = 0,
    quantum: Decimal = CENT,
) -> CommissionResult:
    """Tính hoa hồng và tổng thanh toán.

    * Revenue = round(Base x Rate)
    * VAT_ADD:    VAT = round(Revenue x tax.rate);  Total = Revenue + VAT
    * TAX_DEDUCT: Khấu trừ = round(Revenue x tax.rate); Total = Revenue - Khấu trừ
    * Total còn trừ ``collected_on_behalf`` (thu hộ) và cộng ``consulting_fee``.

    Làm tròn từng bước (như hoá đơn) với ROUND_HALF_UP tới ``quantum`` (0.01 hoặc 1 cho VND chẵn).

    >>> r = calculate_commission(Decimal("88402600"), Decimal("0.18"), DEFAULT_TAX_RULES["EZS"][0])
    >>> (r.revenue, r.vat_amount, r.total_payable)
    (Decimal('15912468.00'), Decimal('1272997.44'), Decimal('17185465.44'))
    """
    base = _to_decimal(commission_base, "commission_base")
    rate = _to_decimal(commission_rate, "commission_rate")
    d = _to_decimal(collected_on_behalf, "collected_on_behalf")
    f = _to_decimal(consulting_fee, "consulting_fee")
    if base < 0:
        raise ValueError("commission_base không được âm")
    if not ZERO <= rate <= 1:
        raise ValueError(f"commission_rate ngoài [0, 1]: {rate} (18% -> Decimal('0.18'))")
    if d < 0 or f < 0:
        raise ValueError("collected_on_behalf / consulting_fee không được âm")

    def q(x: Decimal) -> Decimal:
        return x.quantize(quantum, rounding=ROUND_HALF_UP)

    revenue = q(base * rate)
    tax_amount = q(revenue * tax.rate)
    vat = tax_amount if tax.mode is TaxMode.VAT_ADD else q(ZERO)
    deduction = tax_amount if tax.mode is TaxMode.TAX_DEDUCT else q(ZERO)
    d, f = q(d), q(f)
    return CommissionResult(
        commission_base=base,
        commission_rate=rate,
        revenue=revenue,
        tax_mode=tax.mode,
        tax_rate=tax.rate,
        vat_amount=vat,
        tax_deduction_amount=deduction,
        collected_on_behalf=d,
        consulting_fee=f,
        total_payable=revenue + vat - deduction - d + f,
    )


# =============================================================================
# Orchestration
# =============================================================================
@dataclass(frozen=True)
class ReconciliationResult:
    parse: ParseResult
    nmv: NmvResult
    commission: CommissionResult

    @property
    def issues(self) -> tuple[RowIssue, ...]:
        return tuple(sorted(self.parse.issues + self.nmv.issues, key=lambda i: i.row_number))

    @property
    def has_blocking_errors(self) -> bool:
        """True nếu có lỗi dữ liệu (ERROR) — API nên chặn phát hành BBNT."""
        return any(i.severity is Severity.ERROR for i in self.issues)


def reconcile(
    rows: Sequence[Sequence[Any]],
    config: MappingConfig,
    *,
    commission_rate: Number,
    tax: TaxPolicy,
    period_months: Iterable[YearMonthLike] | None = None,
    collected_on_behalf: Number = 0,
    consulting_fee: Number = 0,
    quantum: Decimal = CENT,
) -> ReconciliationResult:
    """Parse -> NMV -> hoa hồng/thuế trong 1 lần gọi."""
    parsed = parse_sheet(rows, config)
    nmv = compute_nmv(parsed.lines, period_months)
    commission = calculate_commission(nmv.nmv, commission_rate, tax,
                                      collected_on_behalf=collected_on_behalf,
                                      consulting_fee=consulting_fee, quantum=quantum)
    return ReconciliationResult(parsed, nmv, commission)
