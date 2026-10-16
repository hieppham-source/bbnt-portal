"""Đọc Google Sheets bằng OAuth2 access token của người dùng: file order (link do user dán) và Sheet master."""
from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any, Protocol
from urllib.parse import parse_qs, quote, urlparse

_SHEET_PATH = re.compile(r"^/spreadsheets/d/([A-Za-z0-9_-]{20,})")


class SheetAccessError(RuntimeError):
    """Không đọc được sheet (chưa có quyền xem, sai gid, Google lỗi...)."""


class SheetAuthError(SheetAccessError):
    """Access token Google của người dùng hết hạn / bị thu hồi -> cần đăng nhập lại."""


@dataclass(frozen=True)
class SheetRef:
    spreadsheet_id: str
    gid: int | None = None

    @property
    def canonical_url(self) -> str:
        url = f"https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/edit"
        return f"{url}#gid={self.gid}" if self.gid is not None else url


def parse_sheet_url(url: str) -> SheetRef:
    """Nhận link Google Sheets (``/edit#gid=..``, ``/edit?gid=..``, ``/view``...). Raise ValueError nếu sai."""
    u = urlparse(url.strip())
    if u.scheme != "https" or u.netloc != "docs.google.com":
        raise ValueError("Link phải là https://docs.google.com/spreadsheets/d/...")
    m = _SHEET_PATH.match(u.path)
    if not m:
        raise ValueError("Không tìm thấy spreadsheet id trong link")
    gid_raw = parse_qs(u.query).get("gid", [None])[0] or parse_qs(u.fragment).get("gid", [None])[0]
    if gid_raw is not None and not str(gid_raw).isdigit():
        raise ValueError("gid không hợp lệ")
    return SheetRef(m.group(1), int(gid_raw) if gid_raw is not None else None)


class SheetReader(Protocol):
    def read(self, ref: SheetRef, *, sheet_name: str | None = None) -> list[list[Any]]:
        """Trả về ma trận ô (UNFORMATTED_VALUE, ngày dạng serial number)."""
        ...


class GoogleSheetsReader:
    """Google Sheets API v4 bằng **OAuth2 access token của người dùng** (scope ``spreadsheets.readonly``).

    Frontend Next.js xin quyền này lúc đăng nhập Google và gửi token kèm mỗi request (``X-Google-Access-Token``).
    Ai mở được Sheet nào trên Drive thì tool đọc được Sheet đó cho người đó — không dùng service account.
    Token chỉ sống trong request, không lưu, không ghi log.

    Ưu tiên tab theo ``gid`` trong link; không có gid thì dùng ``sheet_name`` của mapping; không có nữa thì tab đầu.
    """

    BASE = "https://sheets.googleapis.com/v4/spreadsheets"

    def __init__(self, access_token: str, *, timeout: float = 30.0, session: Any = None) -> None:
        if not access_token:
            raise ValueError("Thiếu access token Google")
        if session is None:
            import requests

            session = requests.Session()
        session.headers["Authorization"] = f"Bearer {access_token}"
        self._session: Any = session
        self._timeout = timeout

    @classmethod
    def from_access_token(cls, token: str, **kw: Any) -> GoogleSheetsReader:
        return cls(token, **kw)

    def _check(self, resp: Any, what: str = "sheet") -> None:
        if resp.status_code == 401:
            raise SheetAuthError("Phiên Google hết hạn — đăng nhập lại")
        if resp.status_code in (403, 404):
            raise SheetAccessError(f"Tài khoản của bạn không mở được {what} — kiểm tra link và nhờ chủ file "
                                   "share quyền Người xem cho email của bạn")

    def _get(self, url: str, params: dict[str, str]) -> dict[str, Any]:
        resp = self._session.get(url, params=params, timeout=self._timeout)
        self._check(resp)
        if resp.status_code == 400:
            body = resp.text if isinstance(getattr(resp, "text", None), str) else ""
            if "not supported for this document" in body:
                raise SheetAccessError("Link này là file Excel (.xlsx) trên Drive, không phải Google Sheet — mở file "
                                       "› Tệp › Lưu dưới dạng Google Trang tính, rồi dán link của bản mới")
            raise SheetAccessError("Không đọc được tab trong sheet — kiểm tra tên tab (mapping) hoặc gid trong link")
        if resp.status_code >= 400:
            raise SheetAccessError(f"Google Sheets API lỗi {resp.status_code}")
        data: dict[str, Any] = resp.json()
        return data

    def read(self, ref: SheetRef, *, sheet_name: str | None = None) -> list[list[Any]]:
        title = sheet_name
        if ref.gid is not None or title is None:
            meta = self._get(f"{self.BASE}/{ref.spreadsheet_id}", {"fields": "sheets.properties(sheetId,title)"})
            sheets = [s["properties"] for s in meta.get("sheets", [])]
            if not sheets:
                raise SheetAccessError("Spreadsheet không có tab nào")
            if ref.gid is not None:
                match = next((s for s in sheets if s["sheetId"] == ref.gid), None)
                if match is None:
                    raise SheetAccessError(f"Không có tab gid={ref.gid}")
                title = match["title"]
            else:
                title = sheets[0]["title"]
        a1 = "'" + title.replace("'", "''") + "'"
        data = self._get(f"{self.BASE}/{ref.spreadsheet_id}/values/{quote(a1, safe='')}",
                         {"valueRenderOption": "UNFORMATTED_VALUE", "dateTimeRenderOption": "SERIAL_NUMBER",
                          "majorDimension": "ROWS"})
        values: list[list[Any]] = data.get("values", [])
        return values

    def titles(self, ref: SheetRef) -> list[str]:
        meta = self._get(f"{self.BASE}/{ref.spreadsheet_id}", {"fields": "sheets.properties(title)"})
        return [s["properties"]["title"] for s in meta.get("sheets", [])]

    def read_tabs(self, ref: SheetRef, titles: list[str]) -> dict[str, list[list[Any]]]:
        """Đọc nhiều tab trong 1 request (values:batchGet)."""
        if not titles:
            return {}
        params: list[tuple[str, str]] = [("ranges", "'" + t.replace("'", "''") + "'") for t in titles]
        params += [("valueRenderOption", "UNFORMATTED_VALUE"), ("dateTimeRenderOption", "SERIAL_NUMBER"),
                   ("majorDimension", "ROWS")]
        resp = self._session.get(f"{self.BASE}/{ref.spreadsheet_id}/values:batchGet", params=params,
                                 timeout=self._timeout)
        self._check(resp, "Sheet master")
        if resp.status_code >= 400:
            raise SheetAccessError(f"Google Sheets API lỗi {resp.status_code}")
        ranges = resp.json().get("valueRanges", [])
        return {t: r.get("values", []) for t, r in zip(titles, ranges, strict=False)}


def read_xlsx_tabs(path: str, titles: list[str] | None = None) -> dict[str, list[list[Any]]]:
    """Đọc các tab của file Excel (dùng cho CLI / test: Sheet master tải về dạng .xlsx)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = wb.sheetnames if titles is None else [t for t in titles if t in wb.sheetnames]
        out: dict[str, list[list[Any]]] = {}
        for n in names:
            rows = [list(r) for r in wb[n].iter_rows(values_only=True)]
            while rows and all(c is None or (isinstance(c, str) and not c.strip()) for c in rows[-1]):
                rows.pop()      # ô đã định dạng sẵn nhưng trống -> openpyxl vẫn trả về dòng rỗng
            out[n] = rows
        return out
    finally:
        wb.close()


def xlsx_sheet_names(path: str) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
