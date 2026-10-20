"""API stateless: master từ Sheet (giả lập), dán link order -> preview / tải .docx; server không lưu gì."""
from __future__ import annotations

import io
from decimal import Decimal
from typing import Any
from urllib.parse import unquote

import docx
import pytest
from fastapi.testclient import TestClient

from app.api import deps
from app.config import Settings
from app.main import create_app
from app.master.repository import MasterRepository
from app.services.generator import DOCX_MIME, BbntGenerator
from app.services.templates import TemplateRegistry

from .conftest import (
    MASTER_XLSX,
    ORDER_URL_FB,
    ORDER_URL_SHISEIDO,
    FakeSheets,
    order_rows,
    sheet_id,
    vn_clock,
)

USER = {"X-User-Email": "hiep.pham@hannaholala.com"}
ORGANICA = {"channel": "FB", "entity": "EZS", "brand_codes": ["ORGANICA"], "bbnt_suffix": "05",
            "contract_no": "01/2026/HĐDV/EZS-ORGANICA", "campaign_name": "BDAY", "order_sheet_url": ORDER_URL_FB}
SHISEIDO = {"channel": "FB", "entity": "SASSY", "group_code": "SHISEIDO_GRP", "bbnt_suffix": "CAMP",
            "contract_no": "05/2026/HĐDV/SS-SHISEIDO", "campaign_name": "Q3", "order_sheet_url": ORDER_URL_SHISEIDO}


class Env:
    def __init__(self, tables: dict[str, list[list[Any]]], sheets: FakeSheets) -> None:
        self.tables, self.sheets = tables, sheets
        self.clock = vn_clock(2026, 6, 3)
        self.master_reads = 0
        settings = Settings.from_env({"MASTER_FILE": str(MASTER_XLSX), "DEV_AUTH_HEADER": "true"})
        self.app = create_app(settings)

        def source() -> Any:
            self.master_reads += 1
            return self.tables

        self.repo = MasterRepository(source, ttl_seconds=60)
        self.templates = TemplateRegistry(settings.templates_dir)
        o = self.app.dependency_overrides
        o[deps.get_master_repo] = lambda: self.repo
        o[deps.get_sheet_reader] = lambda: self.sheets
        o[deps.get_templates] = lambda: self.templates
        o[deps.get_generator] = lambda: BbntGenerator(self.templates, lambda: self.clock())
        self.client = TestClient(self.app)

    def post(self, path: str, body: dict[str, Any], **params: Any) -> Any:
        return self.client.post(path, json=body, headers=USER, params=params)


@pytest.fixture
def env(master_tables, fake_sheets) -> Env:
    return Env(master_tables, fake_sheets)


def docx_text(content: bytes) -> str:
    d = docx.Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


# =========================================================================== preview
def test_preview_organica(env):
    r = env.post("/api/v1/bbnt/preview", ORGANICA)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["bbnt_number"] == "06.26/BBNT-ORGANICA-EZ-FB-05"
    assert b["file_name"] == "BBNT_06.26_BBNT-ORGANICA-EZ-FB-05.docx"
    assert b["program_name"] == "ORGANICA BDAY 05" and b["issue_date"] == "2026-06-03"
    assert b["contract"]["contract_no"] == "01/2026/HĐDV/EZS-ORGANICA" and b["contract"]["master_row"] == 2
    f = {k: Decimal(str(v)) if isinstance(v, (int, float, str)) and k != "tax_mode" and k != "tax_label"
         and k != "total_payable_words" else v for k, v in b["figures"].items()}
    assert f["gmv_estimated"] == 7_200_000                      # Đã giao + hoàn + huỷ, đúng brand, TOÀN BỘ dòng
    assert f["nmv"] == 4_200_000                                # chỉ Đã giao
    assert f["commission_amount"] == 756_000 and f["vat_amount"] == 60_480
    assert f["tax_deduction_amount"] == 0 and f["total_payable"] == 816_480
    assert b["figures"]["tax_label"] == "VAT 8%"
    assert b["figures"]["total_payable_words"].startswith("Tám trăm mười sáu nghìn bốn trăm tám mươi đồng")
    assert b["other_brand_line_count"] == 1                     # dòng Anessa bị loại
    assert {w["code"] for w in b["warnings"]} == {"RETURNED", "CANCELLED"}
    assert [i["product_name"] for i in b["items"]] == ["Serum Organica", "Kem dưỡng Organica"]


def test_preview_group_sassy(env):
    env.clock = vn_clock(2026, 10, 5)
    b = env.post("/api/v1/bbnt/preview", SHISEIDO).json()
    assert b["bbnt_number"] == "10.26/BBNT-SHISEIDO_GRP-SS-FB-CAMP"
    assert b["brands"] == ["SHISEIDO", "ANESSA"] and b["scope_code"] == "SHISEIDO_GRP"
    assert Decimal(b["figures"]["nmv"]) == 15_000_000
    assert Decimal(b["figures"]["commission_amount"]) == 2_250_000
    assert Decimal(b["figures"]["tax_deduction_amount"]) == 22_500
    assert Decimal(b["figures"]["vat_amount"]) == 0
    assert Decimal(b["figures"]["total_payable"]) == 2_227_500


def test_contract_tax_rate_override_changes_vat(env):
    """Cột "Thuế suất riêng" (tuỳ chọn, tab HOP_DONG) ghi đè % VAT/khấu trừ CHO RIÊNG brand đó — không được
    lấy 1 con số thuế cố định dùng chung cho mọi brand cùng pháp nhân (vd NQ 204/2025/QH15)."""
    h = env.tables["HOP_DONG"]
    h[1][h[0].index("Thuế suất riêng")] = "5%"                 # dòng hợp đồng ORGANICA (EZS, VAT mặc định 8%)
    b = env.post("/api/v1/bbnt/preview", ORGANICA, refresh="true").json()
    assert b["figures"]["tax_rate"] == "0.05"
    assert "HD" not in b["figures"]["tax_label"] or "riêng theo HĐ" in b["figures"]["tax_label"]
    assert "01/2026/HĐDV/EZS-ORGANICA" in b["figures"]["tax_label"]
    assert Decimal(b["figures"]["commission_amount"]) == 756_000           # hoa hồng không đổi
    assert Decimal(b["figures"]["vat_amount"]) == 37_800                   # 756.000 x 5% (thay vì 8%)
    assert Decimal(b["figures"]["total_payable"]) == 793_800


def test_contract_appendix_renders_plhd_clause(env):
    """Hợp đồng có "Số phụ lục HĐ" + "Ngày ký phụ lục" (PLHD, cả 2 cột tuỳ chọn, tab HOP_DONG) -> câu "Căn cứ..."
    trong .docx in thêm "và Phụ lục hợp đồng số ... ký ngày ..." NGAY SAU ngày ký hợp đồng GỐC — không được nhét
    số phụ lục trước số hợp đồng gốc (lỗi người dùng báo trước đây: "...PLHD .../PLHD/... (HĐ gốc ...) ký ngày
    ..." bị ngược)."""
    h = env.tables["HOP_DONG"]
    h[0] = [*h[0], "Số phụ lục HĐ", "Ngày ký phụ lục"]
    h[1] = [*h[1], "PLHD-01/ORGANICA", "31/07/2026"]            # dòng hợp đồng ORGANICA (EZS) — Ngày ký trống
    r = env.post("/api/v1/bbnt/preview", ORGANICA, refresh="true")
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["contract"]["appendix_no"] == "PLHD-01/ORGANICA"
    assert b["contract"]["appendix_signed_on"] == "2026-07-31"

    r2 = env.post("/api/v1/bbnt/generate", ORGANICA, refresh="true")
    assert r2.status_code == 200, r2.text
    text = docx_text(r2.content)
    assert "01/2026/HĐDV/EZS-ORGANICA" in text
    assert ("ký ngày 01 tháng 01 năm 2026 và Phụ lục hợp đồng số PLHD-01/ORGANICA ký ngày "
            "31 tháng 07 năm 2026") in text


# =========================================================================== hợp đồng tự điền (manual_contract)
def test_manual_contract_skips_lookup_and_zeroes_commission(env):
    """``manual_contract=True`` -> bỏ qua tra cứu Sheet master (``contract_no`` không cần có trong tab HOP_DONG),
    dùng thẳng số hợp đồng tự gõ; Bên A/% hoa hồng/ngày ký đều để PLACEHOLDER RÕ RÀNG, hoa hồng/VAT/khấu trừ/tổng
    thanh toán đều tính ra 0 — người làm PHẢI tự điền tay + tự tính tiền trên file .docx đã tải."""
    payload = {**ORGANICA, "contract_no": "CHƯA CÓ TRONG HỆ THỐNG", "manual_contract": True}
    r = env.post("/api/v1/bbnt/preview", payload)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["manual_contract"] is True
    assert b["bbnt_number"] == "06.26/BBNT-ORGANICA-EZ-FB-05"           # số BBNT không phụ thuộc hợp đồng
    assert b["contract"]["contract_no"] == "CHƯA CÓ TRONG HỆ THỐNG"
    assert b["contract"]["party_a_name"] == "(Tự điền sau khi tải BBNT)"
    assert b["contract"]["commission_rate"] == "0"
    assert b["contract"]["signed_on"] == b["issue_date"]                # ngày ký = ngày lập BBNT (placeholder)
    f = b["figures"]
    assert Decimal(f["commission_amount"]) == 0 and Decimal(f["vat_amount"]) == 0
    assert Decimal(f["tax_deduction_amount"]) == 0 and Decimal(f["total_payable"]) == 0
    assert Decimal(f["nmv"]) == 4_200_000                               # NMV vẫn tính bình thường từ đơn hàng

    r2 = env.post("/api/v1/bbnt/generate", payload)
    assert r2.status_code == 200, r2.text                               # không bị MissingValueError
    text = docx_text(r2.content)
    assert "(Tự điền sau khi tải BBNT)" in text
    assert "CHƯA CÓ TRONG HỆ THỐNG" in text
    assert "{{" not in text


def test_manual_contract_group_scope_has_no_brand_code(env):
    """Nghiệm thu nhiều brand cùng group (``scope.is_group``) với ``manual_contract=True`` -> hợp đồng placeholder
    không gắn riêng 1 brand nào (``brand_code=None``), chỉ gắn group — khớp quy tắc hợp đồng gộp group bình
    thường trong :func:`app.master.index.find_contract`."""
    payload = {**SHISEIDO, "manual_contract": True}
    b = env.post("/api/v1/bbnt/preview", payload).json()
    assert b["manual_contract"] is True
    assert b["contract"]["contract_no"] == SHISEIDO["contract_no"]
    assert Decimal(b["figures"]["total_payable"]) == 0
    assert b["brands"] == ["SHISEIDO", "ANESSA"]                        # scope/brand vẫn tách đúng như thường


# =========================================================================== nhiều link order (extra_order_sheet_urls)
SECOND_ORDER_URL = "https://docs.google.com/spreadsheets/d/1OrderFbSecondAbCdEfGhIjKlMnOp/edit"


def test_extra_order_sheet_urls_merges_lines_from_multiple_links(env):
    """Dán thêm ``extra_order_sheet_urls`` -> dòng đơn hàng của link đó được GỘP vào cùng 1 BBNT (vẫn tách
    theo brand/pháp nhân như lúc chọn — chỉ gộp thêm dữ liệu, không tự đổi scope/entity)."""
    env.sheets.sheets[sheet_id(SECOND_ORDER_URL)] = order_rows([
        ["04/05/2026", "FB101", "Đã giao", "Organica", "ORG-01", "Serum Organica", 1, 1_200_000],
    ])
    base = env.post("/api/v1/bbnt/preview", ORGANICA).json()
    merged = env.post("/api/v1/bbnt/preview", {**ORGANICA, "extra_order_sheet_urls": [SECOND_ORDER_URL]}).json()
    assert Decimal(merged["figures"]["nmv"]) == Decimal(base["figures"]["nmv"]) + 1_200_000
    assert Decimal(merged["figures"]["gmv_estimated"]) == Decimal(base["figures"]["gmv_estimated"]) + 1_200_000
    assert merged["order_tabs_combined"] == ["ORDER", "ORDER"]
    assert merged["order_tab"] is None and merged["order_tabs"] == []


def test_breakdown_camp_across_multiple_order_links_for_single_brand(env):
    """1 brand, gộp từ 2 link order RIÊNG BIỆT trở lên (``extra_order_sheet_urls``) -> ``breakdown_mode='camp'``
    tách riêng bảng doanh thu/sản phẩm theo TỪNG LINK (``link_index``), không chỉ theo tab trong CÙNG 1 link
    như :func:`test_breakdown_camp_merges_all_candidate_tabs_of_1_brand` ở ``test_order_files.py`` — đúng tình
    huống dữ liệu 1 brand nằm rải trên nhiều file order khác nhau theo từng đợt chạy/camp."""
    env.sheets.sheets[sheet_id(SECOND_ORDER_URL)] = order_rows([
        ["04/05/2026", "FB101", "Đã giao", "Organica", "ORG-01", "Serum Organica", 1, 1_200_000],
    ])
    base = env.post("/api/v1/bbnt/preview", ORGANICA).json()
    payload = {**ORGANICA, "extra_order_sheet_urls": [SECOND_ORDER_URL], "breakdown_mode": "camp",
              "camp_names": [{"link_index": 0, "tab": "ORDER", "name": "Link chính"},
                             {"link_index": 1, "tab": "ORDER", "name": "Link phụ"}]}
    b = env.post("/api/v1/bbnt/preview", payload).json()
    assert b["breakdown_kind"] == "camp"
    camps = b["brand_breakdown"]
    assert [c["brand_name"] for c in camps] == ["Link chính", "Link phụ"]
    assert [c["brand_code"] for c in camps] == [None, None]
    assert Decimal(camps[0]["nmv"]) == Decimal(base["figures"]["nmv"])
    assert Decimal(camps[1]["nmv"]) == 1_200_000
    assert Decimal(b["figures"]["nmv"]) == Decimal(base["figures"]["nmv"]) + 1_200_000


def test_generate_rejects_split_tab_mode(env):
    """``tab_mode='split'`` chỉ dùng được qua ``/preview-split`` — gọi ``/generate`` thẳng phải báo lỗi."""
    r = env.post("/api/v1/bbnt/generate", {**ORGANICA, "tab_mode": "split", "tab_campaign_names": {"ORDER": "X"}})
    assert r.status_code == 422 and r.json()["code"] == "INVALID_INPUT"


def test_breakdown_mode_rejected_for_single_brand_scope(env):
    """``breakdown_mode`` (tách bảng theo brand/camp) chỉ dùng khi nghiệm thu group/nhiều brand cùng lúc."""
    r = env.post("/api/v1/bbnt/preview", {**ORGANICA, "breakdown_mode": "brand"})
    assert r.status_code == 422


# =========================================================================== generate
def test_generate_returns_docx_and_stores_nothing(env, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    r = env.post("/api/v1/bbnt/generate", ORGANICA)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == DOCX_MIME
    assert unquote(r.headers["x-bbnt-number"]) == "06.26/BBNT-ORGANICA-EZ-FB-05"
    assert r.headers["x-bbnt-total-payable"] == "816480" and r.headers["x-bbnt-warnings"] == "2"
    assert "BBNT_06.26_BBNT-ORGANICA-EZ-FB-05.docx" in r.headers["content-disposition"]
    assert r.headers["cache-control"] == "no-store"
    text = docx_text(r.content)
    for s in ("06.26/BBNT-ORGANICA-EZ-FB-05", "ORGANICA BDAY 05", "01/2026/HĐDV/EZS-ORGANICA",
              "CÔNG TY TNHH ORGANICA (MẪU)", "0071000123456", "816.480", "Serum Organica"):
        assert s in text, s
    assert "{{" not in text
    assert list(tmp_path.iterdir()) == []                       # không ghi file nào ra đĩa


def test_redo_gives_same_number_and_rereads_order(env):
    a = env.post("/api/v1/bbnt/generate", ORGANICA)
    env.sheets.sheets[sheet_id(ORDER_URL_FB)].append(
        ["25/05/2026", "FB007", "Đã giao", "Organica", "ORG-01", "Serum Organica", 1, 1_000_000])
    b = env.post("/api/v1/bbnt/generate", ORGANICA)
    assert a.headers["x-bbnt-number"] == b.headers["x-bbnt-number"]     # làm lại -> cùng số
    assert (a.headers["x-bbnt-total-payable"], b.headers["x-bbnt-total-payable"]) == ("816480", "1010880")
    assert len(env.sheets.calls) == 2                           # file order luôn đọc mới, không cache
    env.clock = vn_clock(2026, 7, 2)                            # lập sang tháng sau -> tiền tố tháng lập đổi
    c = env.post("/api/v1/bbnt/preview", ORGANICA).json()
    assert c["bbnt_number"] == "07.26/BBNT-ORGANICA-EZ-FB-05"


def test_master_edit_picked_up_after_refresh(env):
    assert env.post("/api/v1/bbnt/preview", ORGANICA).json()["contract"]["commission_rate"] == "0.18"
    h = env.tables["HOP_DONG"]
    h[1][h[0].index("% hoa hồng")] = "20%"
    assert env.post("/api/v1/bbnt/preview", ORGANICA).json()["contract"]["commission_rate"] == "0.18"   # cache
    b = env.post("/api/v1/bbnt/preview", ORGANICA, refresh="true").json()
    assert b["contract"]["commission_rate"] == "0.2" and Decimal(b["figures"]["commission_amount"]) == 840_000
    assert env.master_reads == 2


# =========================================================================== errors
@pytest.mark.parametrize("over, status, code", [
    ({"entity": "SASSY"}, 422, "CONTRACT_NOT_FOUND"),
    ({"contract_no": "KHONG-TON-TAI"}, 422, "CONTRACT_NOT_FOUND"),              # hợp đồng tự gõ không khớp
    ({"brand_codes": ["ORGANICA", "ANESSA"]}, 422, "BRANDS_NOT_SAME_GROUP"),
    ({"brand_codes": ["XYZ"]}, 422, "BRAND_NOT_FOUND"),
    ({"order_sheet_url": "https://docs.google.com/spreadsheets/d/1NotSharedAbCdEfGhIjKlMn/edit"}, 502,
     "ORDER_SHEET_UNREADABLE"),
])
def test_business_errors(env, over, status, code):
    r = env.post("/api/v1/bbnt/generate", {**ORGANICA, **over})
    assert r.status_code == status, r.text
    assert r.json()["code"] == code


def test_no_delivered_orders(env):
    env.sheets.sheets[sheet_id(ORDER_URL_FB)] = order_rows([
        ["03/05/2026", "FB001", "Đã hủy", "Organica", "ORG-01", "Serum", 1, 100_000]])
    r = env.post("/api/v1/bbnt/generate", ORGANICA)
    assert r.status_code == 422 and r.json()["code"] == "NO_DELIVERED_ORDERS"


def test_sheet_data_errors_listed_by_row(env):
    env.sheets.sheets[sheet_id(ORDER_URL_FB)] = order_rows([
        ["03/05/2026", "FB001", "Đã giao", "Organica", "ORG-01", "Serum", 1, "abc"]])
    r = env.post("/api/v1/bbnt/generate", ORGANICA)
    body = r.json()
    assert r.status_code == 422 and body["code"] == "SHEET_DATA_ERRORS"
    assert body["details"][0]["row_number"] == 3


def test_mapping_mismatch(env):
    env.sheets.sheets[sheet_id(ORDER_URL_FB)] = [["Mã đơn", "Tiền"], ["X", 1]]
    r = env.post("/api/v1/bbnt/generate", ORGANICA)
    assert r.status_code == 422 and r.json()["code"] == "MAPPING_MISMATCH"


@pytest.mark.parametrize("bad", [
    {"brand_codes": None},
    {"group_code": "SHISEIDO_GRP"},                                          # cả brand lẫn group
    {"order_sheet_url": "https://example.com/sheet"},
    {"bbnt_suffix": ""},
    {"contract_no": ""},
    {"collected_on_behalf": -1},
    {"unknown": 1},
    {"extra_order_sheet_urls": [f"https://docs.google.com/spreadsheets/d/1ExtraLink{i:02d}AbCdEfGhIjKlMnOp/edit"
                               for i in range(10)]},                      # tối đa 9 (10 link cả order_sheet_url)
])
def test_payload_validation(env, bad):
    body = {k: v for k, v in {**ORGANICA, **bad}.items() if v is not None}
    assert env.post("/api/v1/bbnt/preview", body).status_code == 422


def test_auth_required(env):
    assert env.client.post("/api/v1/bbnt/generate", json=ORGANICA).status_code == 401
    assert env.client.get("/api/v1/master").status_code == 401


# =========================================================================== master endpoints
def test_master_options(env):
    b = env.client.get("/api/v1/master", headers=USER).json()
    assert b["channels"] == ["FB", "TTS"]
    assert {e["code"] for e in b["entities"]} == {"EZS", "SASSY"}
    assert {(g["channel"], g["code"]) for g in b["groups"]} == {("FB", "SHISEIDO_GRP")}
    assert next(g for g in b["groups"])["brand_codes"] == ["SHISEIDO", "ANESSA"]
    assert len(b["contracts"]) == 2 and b["mappings"][0]["is_default"] is True
    assert b["error_count"] == 0


def test_master_check_lists_issues(env):
    env.tables["BRAND"].append(["FB", "BAD CODE", "x"])
    b = env.client.get("/api/v1/master/check", headers=USER).json()
    assert b["ok"] is False
    assert b["issues"][0]["tab"] == "BRAND" and b["issues"][0]["row"] == 6


def test_readyz(env):
    r = env.client.get("/readyz")
    assert r.status_code == 200 and r.json() == {"status": "ok", "master": "ok", "templates": "ok"}


def test_app_uses_its_own_settings_without_overrides(monkeypatch):
    """create_app(settings) phải dùng đúng settings truyền vào (không đọc lại biến môi trường)."""
    monkeypatch.delenv("MASTER_SHEET_URL", raising=False)
    monkeypatch.delenv("MASTER_FILE", raising=False)
    settings = Settings.from_env({"MASTER_FILE": str(MASTER_XLSX), "DEV_AUTH_HEADER": "true",
                                  "CORS_ALLOWED_ORIGINS": "http://localhost:3000"})
    client = TestClient(create_app(settings))
    assert client.get("/api/v1/master", headers=USER).json()["channels"] == ["FB", "TTS"]
    assert client.get("/readyz").json()["master"] == "ok"
    pre = client.options("/api/v1/master", headers={"Origin": "http://localhost:3000",
                                                    "Access-Control-Request-Method": "GET",
                                                    "Access-Control-Request-Headers": "x-user-email"})
    assert pre.status_code == 200                       # local dev: frontend gửi X-User-Email được


def test_order_without_brand_column_keeps_all_lines(env):
    """BRAND_NAME tuỳ chọn: file order không có cột Brand -> không lọc brand (trước đây loại hết dòng)."""
    header = ["Ngày đặt hàng", "Mã đơn hàng", "Trạng thái", "SKU", "Tên sản phẩm", "Số lượng",
              "Doanh thu sau voucher"]
    env.sheets.sheets[sheet_id(ORDER_URL_FB)] = [["BÁO CÁO"], header,
                                                 ["03/05/2026", "FB001", "Đã giao", "ORG-01", "Serum", 2, 1_000_000]]
    r = env.post("/api/v1/bbnt/preview", ORGANICA)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["other_brand_line_count"] == 0 and Decimal(str(b["figures"]["nmv"])) == 1_000_000
