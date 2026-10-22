"""File order "lộn xộn" của brand: nhiều tab, .xlsx trên Drive, header lệch dòng / lệch tên, dòng công thức trống;
và ngày ký hợp đồng điền vào BBNT."""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from typing import Any

import docx
import pytest
from openpyxl import Workbook

from app.core.calculation import FieldSpec, MappingConfig, MappingError, TargetField, locate_header, parse_sheet
from app.master.index import Scope
from app.master.models import Brand
from app.services import sheets as sheets_mod
from app.services.order_tabs import detect_brand_tabs
from app.services.sheets import DriveScopeError, OrderWorkbook, SheetAccessError, parse_sheet_url, xlsx_workbook

from .conftest import ORDER_URL_FB, sheet_id
from .test_api import ORGANICA, Env, docx_text

T = TargetField
HEADER = ["Ngày đặt hàng", "Tên người nhận", "Mã đơn hàng", "Hãng", "Mã sản phẩm *", "Tên sản phẩm", "Số lượng",
          "Total GMV", "GMV SAU VOUCHER (IFANY)", "BRAND INPUT\nTình trạng lên đơn", "BRAND INPUT\nNote"]


def tab(brand: str, lines: list[list[Any]], filler: int = 3) -> list[list[Any]]:
    """Kiểu file brand: dòng 1 là tổng, dòng 2 header, cuối có dòng kẻ sẵn công thức (GMV = 0)."""
    total = [None] * 8 + [sum(r[5] for r in lines)]
    rows = [[d, "Khách", code, brand, sku, name, qty, amt, amt, st, None] for d, code, sku, name, qty, amt, st in lines]
    return [total, HEADER, *rows, *([[None] * 8 + [0]] * filler)]


ORG = tab("ORGANICA", [["03/05/2026", "H1", "ORG-01", "Serum", 2, 1_000_000, "ĐÃ GIAO"],
                       ["10/05/2026", "H2", "ORG-02", "Kem", 1, 2_500_000, "ĐÃ GIAO"],
                       ["11/05/2026", "H3", "ORG-02", "Kem", 1, 2_500_000, "YÊU CẦU HỦY"]])
ANESSA = tab("ANESSA", [["05/05/2026", "H9", "ANS-01", "Kem chống nắng", 1, 9_000_000, "ĐÃ GIAO"]])
SUMMARY = [[None, "BIÊN BẢN NGHIỆM THU ORGANICA"], [None, "GMV from HRV", 3_500_000]]


@pytest.fixture
def env(master_tables, fake_sheets) -> Env:
    return Env(master_tables, fake_sheets)


def put(env: Env, tabs: dict[str, list[list[Any]]] | list[list[Any]]) -> None:
    env.sheets.sheets[sheet_id(ORDER_URL_FB)] = tabs


# ----------------------------------------------------------------------------- chọn tab
def test_multi_tab_file_picks_tab_of_selected_brand(env):
    put(env, {"ANESSA ORDER MIX 05.26": ANESSA, "ORGANICA ORDER MIX 05.26": ORG, "BBNT ORGANICA": SUMMARY})
    r = env.post("/api/v1/bbnt/preview", ORGANICA)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["order_tab"] == "ORGANICA ORDER MIX 05.26" and b["order_header_row"] == 2
    assert b["order_tabs"] == ["ANESSA ORDER MIX 05.26", "ORGANICA ORDER MIX 05.26"]     # tab BBNT bị loại
    assert Decimal(b["figures"]["nmv"]) == 3_500_000 and Decimal(b["figures"]["gmv_estimated"]) == 6_000_000
    assert {w["code"] for w in b["warnings"]} == {"CANCELLED"}       # 'YÊU CẦU HỦY' = huỷ; dòng công thức bị bỏ qua


NO_GID = {**ORGANICA, "order_sheet_url": ORDER_URL_FB.split("#")[0]}


def test_ambiguous_tabs_ask_user_then_use_choice(env):
    put(env, {"ORDER 1": ORG, "ORDER 2": ORG})
    r = env.post("/api/v1/bbnt/preview", NO_GID)
    assert r.status_code == 409 and r.json()["code"] == "ORDER_TAB_AMBIGUOUS"
    assert r.json()["details"]["tabs"] == ["ORDER 1", "ORDER 2"]
    r = env.post("/api/v1/bbnt/preview", {**NO_GID, "order_tab": "ORDER 2 "})
    assert r.status_code == 200 and r.json()["order_tab"] == "ORDER 2"


def test_campaign_name_breaks_tie_between_tabs_of_same_brand(env):
    put(env, {"ORGANICA ORDER FOOD 05.26": ANESSA, "ORGANICA ORDER BDAY 05.26": ORG})
    assert env.post("/api/v1/bbnt/preview", ORGANICA).json()["order_tab"] == "ORGANICA ORDER BDAY 05.26"


def test_gid_does_not_resolve_a_pure_tie_even_between_just_2_tabs(env):
    """2 tab hoàn toàn không có tín hiệu nội dung nào (không khớp brand/chương trình) -> gid (tab đang mở khi
    copy link) không còn đủ để tự chọn 1 tab, kể cả chỉ có 2 tab — phải báo ORDER_TAB_AMBIGUOUS. Trước đây gid
    từng tự chọn thẳng 1 tab trong trường hợp này, nhưng đó chính là kiểu lỗi âm thầm bỏ sót camp còn lại khi
    file chỉ có đúng 2 tab/camp (xem test_gid_does_not_silently_pick_1_tab_out_of_a_multi_camp_file bên dưới
    cho 3 tab) — gid chỉ còn tác dụng phá thế hoà khi các tab ĐÃ có tín hiệu nội dung thật (xem test kế tiếp)."""
    put(env, {"ORDER 1": ORG, "ORDER 2": ORG})
    r = env.post("/api/v1/bbnt/preview", {**ORGANICA, "order_sheet_url": ORDER_URL_FB.replace("gid=0", "gid=1")})
    assert r.status_code == 409 and r.json()["code"] == "ORDER_TAB_AMBIGUOUS"


def test_gid_still_breaks_a_tie_when_tabs_already_share_real_content_signal(env):
    """2 tab CÙNG khớp brand + tên chương trình ngang điểm nhau (vd 2 bản ghi/sửa đổi trùng tên) -> gid (tab
    đang mở khi copy link) vẫn hợp lý để chọn giúp 1 trong 2, không cần hỏi lại."""
    put(env, {"ORGANICA BDAY A": ORG, "ORGANICA BDAY B": ORG})
    r = env.post("/api/v1/bbnt/preview", {**ORGANICA, "order_sheet_url": ORDER_URL_FB.replace("gid=0", "gid=1")})
    assert r.status_code == 200 and r.json()["order_tab"] == "ORGANICA BDAY B"


def test_gid_does_not_silently_pick_1_tab_out_of_a_multi_camp_file(env):
    """3+ tab hoàn toàn không tên brand/chương trình nào (vd 1 brand chạy nhiều camp: "CAMPMIX7/8/9", mỗi tab 1
    đợt) — link order đang mở đúng 1 trong số đó (gid) KHÔNG được tự chọn mỗi tab đó rồi âm thầm bỏ qua 2 camp
    còn lại; phải báo ORDER_TAB_AMBIGUOUS để người dùng gộp qua breakdown_mode='camp' (xem
    test_breakdown_camp_merges_all_candidate_tabs_of_1_brand)."""
    put(env, {"CAMPMIX7": ORG, "CAMPMIX8": ORG, "CAMPMIX9": ORG})
    url = ORDER_URL_FB.replace("gid=0", "gid=2")  # đang mở đúng tab CAMPMIX9 khi copy link
    r = env.post("/api/v1/bbnt/preview", {**ORGANICA, "order_sheet_url": url})
    assert r.status_code == 409 and r.json()["code"] == "ORDER_TAB_AMBIGUOUS"
    assert set(r.json()["details"]["tabs"]) == {"CAMPMIX7", "CAMPMIX8", "CAMPMIX9"}
    # vẫn chọn thẳng được nếu người dùng tự xác nhận đúng 1 tab đó
    r = env.post("/api/v1/bbnt/preview", {**ORGANICA, "order_sheet_url": url, "order_tab": "CAMPMIX9"})
    assert r.status_code == 200 and r.json()["order_tab"] == "CAMPMIX9"


def test_unknown_requested_tab_and_no_matching_tab(env):
    put(env, {"ORGANICA": ORG})
    r = env.post("/api/v1/bbnt/preview", {**ORGANICA, "order_tab": "Không có"})
    assert r.status_code == 422 and r.json()["code"] == "ORDER_TAB_NOT_FOUND"
    put(env, {"BBNT": SUMMARY, "Ghi chú": [["abc"]]})
    r = env.post("/api/v1/bbnt/preview", ORGANICA)
    assert r.json()["code"] == "MAPPING_MISMATCH" and "đã thử 2 tab" in r.json()["message"]


# ----------------------------------------------------------------------------- tab_mode split / breakdown_mode=camp
def test_breakdown_camp_merges_all_candidate_tabs_of_1_brand(env):
    """Link có nhiều tab mơ hồ (nhiều camp cùng 1 brand) -> ``breakdown_mode='camp'`` gộp TẤT CẢ tab lại thành
    1 BBNT DUY NHẤT (1 số, 1 hoa hồng chung) thay vì báo ``ORDER_TAB_AMBIGUOUS``, nhưng tách riêng bảng doanh
    thu + danh sách sản phẩm theo từng camp (đặt tên qua ``camp_names``) — thay cho ``tab_mode='combine'`` cũ,
    nay đã gộp chung vào 1 luồng duy nhất với tính năng tách bảng (xem docstring ``app.services.generator``)."""
    put(env, {"ORDER 1": ORG, "ORDER 2": ORG})
    payload = {**NO_GID, "breakdown_mode": "camp",
              "camp_names": [{"link_index": 0, "tab": "ORDER 1", "name": "CAMP1"},
                             {"link_index": 0, "tab": "ORDER 2", "name": "CAMP2"}]}
    b = env.post("/api/v1/bbnt/preview", payload).json()
    assert Decimal(b["figures"]["nmv"]) == 7_000_000 and Decimal(b["figures"]["gmv_estimated"]) == 12_000_000
    assert b["order_tabs_combined"] == ["ORDER 1", "ORDER 2"]
    assert b["order_tab"] is None and b["order_tabs"] == []
    assert b["breakdown_kind"] == "camp"
    camps = b["brand_breakdown"]
    assert [c["brand_name"] for c in camps] == ["CAMP1", "CAMP2"]
    assert [c["brand_code"] for c in camps] == [None, None]
    assert [c["order_tabs"] for c in camps] == [["ORDER 1"], ["ORDER 2"]]
    assert all(Decimal(c["nmv"]) == 3_500_000 and Decimal(c["gmv_estimated"]) == 6_000_000 for c in camps)


def test_breakdown_camp_requires_names_for_every_tab(env):
    """``camp_names`` BẮT BUỘC đặt tên riêng cho TỪNG camp — không truyền gì cả (case Milktouch: tab
    "CAMPMIX7/8/9") báo lỗi ``TAB_CAMPAIGN_NAMES_REQUIRED`` thay vì tự lấy theo tên tab/sheet."""
    put(env, {"CAMPMIX7": ORG, "CAMPMIX8": ORG, "CAMPMIX9": ORG})
    r = env.post("/api/v1/bbnt/preview", {**NO_GID, "breakdown_mode": "camp"})
    assert r.status_code == 409 and r.json()["code"] == "TAB_CAMPAIGN_NAMES_REQUIRED"

    r = env.post("/api/v1/bbnt/preview", {**NO_GID, "breakdown_mode": "camp", "camp_names": [
        {"link_index": 0, "tab": "CAMPMIX7", "name": "Camp 7"},
        {"link_index": 0, "tab": "CAMPMIX8", "name": "Camp 8"},
        {"link_index": 0, "tab": "CAMPMIX9", "name": "Camp 9"},
    ]})
    assert r.status_code == 200, r.text
    b = r.json()
    camps = {c["brand_name"]: c for c in b["brand_breakdown"]}
    assert set(camps) == {"Camp 7", "Camp 8", "Camp 9"}
    assert camps["Camp 7"]["order_tabs"] == ["CAMPMIX7"]


def test_preview_split_returns_one_bbnt_per_tab_with_r_suffix(env):
    """``tab_mode='split'`` (chỉ qua ``/preview-split``) -> mỗi tab ứng viên 1 BBNT riêng, tên chương trình lấy
    từ ``tab_campaign_names`` (bắt buộc khác nhau từng tab), số BBNT giống hệt nhau trừ hậu tố ``-R{n}`` theo
    thứ tự tab (số BBNT không phân biệt theo tên chương trình)."""
    put(env, {"ORDER 1": ORG, "ORDER 2": ORG})
    payload = {**NO_GID, "tab_mode": "split", "tab_campaign_names": {"ORDER 1": "BDAY1", "ORDER 2": "BDAY2"}}
    r = env.post("/api/v1/bbnt/preview-split", payload)
    assert r.status_code == 200, r.text
    results = r.json()["results"]
    assert len(results) == 2
    assert results[0]["order_tab"] == "ORDER 1" and results[0]["campaign_name"] == "BDAY1"
    assert results[1]["order_tab"] == "ORDER 2" and results[1]["campaign_name"] == "BDAY2"
    base = "06.26/BBNT-ORGANICA-EZ-FB-05"
    assert results[0]["bbnt_number"] == f"{base}-R1" and results[1]["bbnt_number"] == f"{base}-R2"
    assert Decimal(results[0]["figures"]["nmv"]) == 3_500_000 and Decimal(results[1]["figures"]["nmv"]) == 3_500_000
    assert results[0]["order_tabs_combined"] == [] and results[1]["order_tabs_combined"] == []


def test_preview_split_requires_distinct_campaign_names_per_tab(env):
    put(env, {"ORDER 1": ORG, "ORDER 2": ORG})
    r = env.post("/api/v1/bbnt/preview-split",
                {**NO_GID, "tab_mode": "split", "tab_campaign_names": {"ORDER 1": "BDAY1"}})
    assert r.status_code == 409 and r.json()["code"] == "TAB_CAMPAIGN_NAMES_REQUIRED"
    assert r.json()["details"]["tabs"] == ["ORDER 2"]


def test_preview_split_rejects_duplicate_campaign_names(env):
    put(env, {"ORDER 1": ORG, "ORDER 2": ORG})
    r = env.post("/api/v1/bbnt/preview-split",
                {**NO_GID, "tab_mode": "split", "tab_campaign_names": {"ORDER 1": "SAME", "ORDER 2": "same "}})
    assert r.status_code == 422 and r.json()["code"] == "TAB_CAMPAIGN_NAMES_DUPLICATE"


def test_preview_split_rejects_multiple_order_links(env):
    put(env, {"ORDER 1": ORG, "ORDER 2": ORG})
    second_url = "https://docs.google.com/spreadsheets/d/1SecondOrderAbCdEfGhIjKlMnOp/edit"
    env.sheets.sheets[sheet_id(second_url)] = ORG
    r = env.post("/api/v1/bbnt/preview-split",
                {**NO_GID, "tab_mode": "split", "extra_order_sheet_urls": [second_url],
                 "tab_campaign_names": {"ORDER 1": "A", "ORDER 2": "B"}})
    assert r.status_code == 422 and r.json()["code"] == "TAB_MODE_NOT_SUPPORTED"


def test_generate_fills_contract_signed_date(env):
    r = env.post("/api/v1/bbnt/preview", ORGANICA).json()
    assert r["contract"]["signed_on_source"] == "Hiệu lực từ"       # cột Ngày ký trống
    assert r["contract"]["signed_on"] == r["contract"]["valid_from"]
    h = env.tables["HOP_DONG"]
    col = h[0].index("Ngày ký")
    for row in h[1:]:
        row.extend([None] * (col + 1 - len(row)))
        row[col] = "25/12/2025"
    env.repo.get(refresh=True)
    b = env.post("/api/v1/bbnt/preview", ORGANICA).json()
    assert b["contract"]["signed_on"] == "2025-12-25" and b["contract"]["signed_on_source"] == "Ngày ký"
    text = docx_text(env.post("/api/v1/bbnt/generate", ORGANICA).content)
    assert "ký ngày 25 tháng 12 năm 2025" in text and "XX tháng XX" not in text


# ----------------------------------------------------------------------------- 1 tab/brand (group, multi-brand)
def _brand(code: str, group_code: str | None = "KOSE") -> Brand:
    return Brand(channel="FB", code=code, name=code, group_code=group_code,
                group_name="KOSE GROUP" if group_code else None, active=True, row=1)


KOSE_GROUP = Scope("FB", (_brand("ANESSA"), _brand("ELIXIR"), _brand("DPROGRAM")), "KOSE", "KOSE GROUP")
_ROW = ["03/05/2026", "H1", "SKU-01", "Sản phẩm", 1, 100_000, "ĐÃ GIAO"]


def workbook(tabs: dict[str, list[list[Any]]]) -> OrderWorkbook:
    return OrderWorkbook("xlsx", list(tabs), lambda t: tabs[t])


def test_detect_brand_tabs_matches_each_brand_to_its_own_tab():
    wb = workbook({"ANESSA": tab("ANESSA", [_ROW]), "ELIXIR": tab("ELIXIR", [_ROW]),
                   "DPROGRAM": tab("DPROGRAM", [_ROW])})
    assert detect_brand_tabs(wb, CFG, scope=KOSE_GROUP) == {"ANESSA": ["ANESSA"], "ELIXIR": ["ELIXIR"],
                                                            "DPROGRAM": ["DPROGRAM"]}


def test_detect_brand_tabs_skips_brand_with_no_matching_tab():
    wb = workbook({"ANESSA": tab("ANESSA", [_ROW]), "ELIXIR": tab("ELIXIR", [_ROW])})
    # DPROGRAM không có tab nào trong file order kỳ này -> bỏ qua, KHÔNG lỗi, KHÔNG có trong map trả về
    assert detect_brand_tabs(wb, CFG, scope=KOSE_GROUP) == {"ANESSA": ["ANESSA"], "ELIXIR": ["ELIXIR"]}


def test_detect_brand_tabs_none_for_single_brand_scope():
    single = Scope("FB", (_brand("ORGANICA", group_code=None),), None, None)
    wb = workbook({"ORGANICA": tab("ORGANICA", [_ROW])})
    assert detect_brand_tabs(wb, CFG, scope=single) is None


def test_detect_brand_tabs_none_when_fewer_than_2_tabs_match_mapping():
    wb = workbook({"ANESSA": tab("ANESSA", [_ROW]), "Ghi chú": [["abc"]]})
    assert detect_brand_tabs(wb, CFG, scope=KOSE_GROUP) is None


def test_detect_brand_tabs_merges_multiple_tabs_of_the_same_brand():
    """1 brand tách nhiều tab theo loại sản phẩm (vd "BIO CORNER ORDER MIX" + "BIO CORNER ORDER FOOD" đều là
    brand BIO CORNER) -> gộp cả 2 tab vào danh sách của brand đó, KHÔNG coi là mơ hồ."""
    wb = workbook({"ANESSA MIX": tab("ANESSA", [_ROW]), "ANESSA FOOD": tab("ANESSA", [_ROW]),
                   "ELIXIR": tab("ELIXIR", [_ROW])})
    result = detect_brand_tabs(wb, CFG, scope=KOSE_GROUP)
    assert result is not None
    assert set(result["ANESSA"]) == {"ANESSA MIX", "ANESSA FOOD"} and result["ELIXIR"] == ["ELIXIR"]


def test_detect_brand_tabs_none_when_tabs_dont_name_any_brand():
    # cả 2 tab đều đọc được theo mapping nhưng tên tab không khớp brand nào trong group -> không phải kiểu
    # "1 tab/brand", trả None để dùng lại choose_order_tab (chấm điểm theo tên/chiến dịch) như file 1 tab chung
    wb = workbook({"ORDER 1": tab("ANESSA", [_ROW]), "ORDER 2": tab("ELIXIR", [_ROW])})
    assert detect_brand_tabs(wb, CFG, scope=KOSE_GROUP) is None


def test_detect_brand_tabs_ignores_tab_whose_name_matches_2_brands():
    # tên tab khớp CẢ 2 brand (vd "ANESSA ELIXIR ORDER") -> mơ hồ cho riêng tab đó, bỏ qua không tính candidate
    # của brand nào — ELIXIR vẫn còn 1 tab riêng khớp rõ nên vẫn đủ điều kiện phát hiện (ANESSA bị bỏ qua)
    wb = workbook({"ANESSA ELIXIR ORDER": tab("ANESSA", [_ROW]), "ELIXIR ORDER 2": tab("ELIXIR", [_ROW]),
                   "DPROGRAM": tab("DPROGRAM", [_ROW])})
    result = detect_brand_tabs(wb, CFG, scope=KOSE_GROUP)
    assert result is not None
    assert "ANESSA" not in result
    assert result["ELIXIR"] == ["ELIXIR ORDER 2"] and result["DPROGRAM"] == ["DPROGRAM"]


# ----------------------------------------------------------------------------- header
CFG = MappingConfig(fields=(
    FieldSpec(T.ORDER_CODE, "Mã đơn hàng", required=True), FieldSpec(T.ORDER_STATUS, "Trạng thái", required=True),
    FieldSpec(T.SKU, "SKU"), FieldSpec(T.GMV_AFTER_VOUCHER, "Doanh thu sau voucher", required=True),
    FieldSpec(T.ORDER_DATE, "Ngày đặt hàng")), header_row=2, data_start_row=3)


def test_default_aliases_and_contains_match():
    _, cols = locate_header([[], HEADER], CFG)
    assert cols[T.ORDER_STATUS] == HEADER.index("BRAND INPUT\nTình trạng lên đơn")
    assert cols[T.GMV_AFTER_VOUCHER] == HEADER.index("GMV SAU VOUCHER (IFANY)")
    assert cols[T.SKU] == HEADER.index("Mã sản phẩm *")


def test_mapping_aliases_with_pipe_take_priority():
    cfg = MappingConfig(fields=(FieldSpec(T.ORDER_CODE, "Order No | Mã đơn hàng"),
                                FieldSpec(T.ORDER_STATUS, "Ghi chú trạng thái | Note"),
                                FieldSpec(T.GMV, "Total GMV")))
    _, cols = locate_header([HEADER], cfg)
    assert cols[T.ORDER_STATUS] == HEADER.index("BRAND INPUT\nNote")


def test_header_row_detected_when_shifted():
    rows = [["BÁO CÁO"], [], ["Tháng 05"], HEADER, ["01/05/2026", "", "H1", "X", "S", "SP", 1, 5, 5, "Đã giao"]]
    hr, _ = locate_header(rows, CFG)
    assert hr == 4
    res = parse_sheet(rows, CFG)
    assert res.header_row == 4 and [ln.row_number for ln in res.lines] == [5]


def test_ambiguous_contains_is_not_guessed():
    cfg = MappingConfig(fields=(FieldSpec(T.ORDER_CODE, "Mã đơn hàng"), FieldSpec(T.ORDER_STATUS, "Status"),
                                FieldSpec(T.GMV, "Doanh số")))
    with pytest.raises(MappingError):
        locate_header([["Mã đơn hàng", "Status cũ", "Status mới", "Doanh số"]], cfg)


# ----------------------------------------------------------------------------- .xlsx & Drive
def xlsx_bytes(tabs: dict[str, list[list[Any]]], hidden: tuple[str, ...] = ()) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tabs.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for _ in range(50):          # ô định dạng sẵn nhưng trống
            ws.append([None])
        if name in hidden:
            ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_workbook_lists_visible_tabs_and_trims():
    org = [[datetime(2026, 5, 3, 10, 0), *r[1:]] if i == 2 else r for i, r in enumerate(ORG)]
    book = xlsx_workbook(xlsx_bytes({"ORGANICA": org, "Ẩn": ANESSA}, hidden=("Ẩn",)))
    assert book.kind == "xlsx" and book.tabs == ["ORGANICA"]
    rows = book.read("ORGANICA")
    assert rows[1] == HEADER and isinstance(rows[2][0], datetime)
    assert len(book.heads(["ORGANICA"])["ORGANICA"]) == 8
    with pytest.raises(SheetAccessError, match="Excel"):
        xlsx_workbook(b"not a zip")


class Resp:
    def __init__(self, status: int, payload: Any = None, *, text: str = "", content: bytes = b"") -> None:
        self.status_code, self._payload, self.text, self.content = status, payload or {}, text, content

    def json(self) -> Any:
        return self._payload


class Session:
    """Giả requests.Session: route theo URL + alt=media."""

    def __init__(self, routes: dict[str, Resp]) -> None:
        self.routes, self.headers, self.calls = routes, {}, []

    def get(self, url: str, params: Any = None, timeout: float = 0) -> Resp:
        p = dict(params) if isinstance(params, dict) else dict(params or [])
        key = ("drive-media" if p.get("alt") == "media" else "drive-meta") if "/drive/" in url else (
            "batch" if url.endswith(":batchGet") else "values" if "/values/" in url else "meta")
        self.calls.append(key)
        return self.routes[key]


NOT_SUPPORTED = Resp(400, text='{"error":{"message":"This operation is not supported for this document"}}')
XLSX_META = Resp(200, {"name": "BÁO CÁO.xlsx", "mimeType": sheets_mod.XLSX_MIME, "size": "1000"})


def reader(routes: dict[str, Resp]) -> sheets_mod.GoogleSheetsReader:
    return sheets_mod.GoogleSheetsReader("tok", session=Session(routes))


def test_google_sheet_open_uses_batch_heads_and_gid():
    r = reader({"meta": Resp(200, {"sheets": [{"properties": {"sheetId": 0, "title": "A"}},
                                              {"properties": {"sheetId": 7, "title": "B"}},
                                              {"properties": {"sheetId": 9, "title": "Ẩn", "hidden": True}}]}),
                "batch": Resp(200, {"valueRanges": [{"values": [["x"]]}, {"values": [["y"]]}]}),
                "values": Resp(200, {"values": [["full"]]})})
    book = r.open(parse_sheet_url("https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUv/edit#gid=7"))
    assert book.kind == "google_sheet" and book.tabs == ["A", "B"] and book.gid_tab == "B"
    assert book.heads(["A", "B"]) == {"A": [["x"]], "B": [["y"]]}
    assert book.read("B") == [["full"]] and book.read("B") == [["full"]]
    assert r._session.calls.count("values") == 1               # nhớ trong request


def test_xlsx_on_drive_is_downloaded_and_read():
    data = xlsx_bytes({"ORGANICA ORDER": ORG})
    r = reader({"meta": NOT_SUPPORTED, "drive-meta": XLSX_META, "drive-media": Resp(200, content=data)})
    book = r.open(parse_sheet_url("https://docs.google.com/spreadsheets/d/1vSkabkV_4w7DlIN9fcBLUNCVsWEdai2w/edit"))
    assert book.kind == "xlsx" and book.tabs == ["ORGANICA ORDER"] and book.read("ORGANICA ORDER")[1] == HEADER


@pytest.mark.parametrize("resp, exc, text", [
    (Resp(403, text='{"error":{"message":"Request had insufficient authentication scopes."}}'),
     DriveScopeError, "đăng nhập lại"),
    (Resp(403, text='{"error":{"reason":"accessNotConfigured","message":"Drive API has not been used"}}'),
     SheetAccessError, "Google Drive API chưa bật"),
    (Resp(404), SheetAccessError, "share quyền"),
])
def test_drive_errors(resp, exc, text):
    r = reader({"meta": NOT_SUPPORTED, "drive-meta": resp})
    with pytest.raises(exc, match=text):
        r.open(parse_sheet_url("https://docs.google.com/spreadsheets/d/1vSkabkV_4w7DlIN9fcBLUNCVsWEdai2w/edit"))


def test_drive_non_excel_file_rejected():
    r = reader({"meta": NOT_SUPPORTED, "drive-meta": Resp(200, {"name": "a.pdf", "mimeType": "application/pdf"})})
    with pytest.raises(SheetAccessError, match="không phải Google Sheet hay Excel"):
        r.open(parse_sheet_url("https://drive.google.com/file/d/1vSkabkV_4w7DlIN9fcBLUNCVsWEdai2w/view"))


def test_api_maps_drive_scope_error(env, monkeypatch):
    def boom(ref: Any) -> Any:
        raise DriveScopeError("cần cấp quyền đọc Drive")
    monkeypatch.setattr(env.sheets, "open", boom)
    r = env.post("/api/v1/bbnt/preview", ORGANICA)
    assert r.status_code == 403 and r.json()["code"] == "GOOGLE_DRIVE_SCOPE_MISSING"


@pytest.mark.parametrize("url, file_id", [
    ("https://drive.google.com/file/d/1vSkabkV_4w7DlIN9fcBLUNCVsWEdai2w/view?usp=drivesdk",
     "1vSkabkV_4w7DlIN9fcBLUNCVsWEdai2w"),
    ("https://drive.google.com/open?id=1vSkabkV_4w7DlIN9fcBLUNCVsWEdai2w", "1vSkabkV_4w7DlIN9fcBLUNCVsWEdai2w"),
])
def test_parse_drive_links(url, file_id):
    ref = parse_sheet_url(url)
    assert ref.spreadsheet_id == file_id and ref.gid is None


def test_docx_still_valid(env):
    put(env, {"ORGANICA ORDER": ORG})
    content = env.post("/api/v1/bbnt/generate", ORGANICA).content
    assert docx.Document(io.BytesIO(content)).paragraphs
