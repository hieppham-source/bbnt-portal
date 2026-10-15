"""Tests docx engine: số thành chữ, giữ định dạng, bảng động, guard ngày pháp lý, tagger."""
from __future__ import annotations

import io
import re
import zipfile
from decimal import Decimal
from pathlib import Path

import docx
import pytest
from lxml import etree

from app.docx import filler as filler_mod
from app.docx.filler import (
    BbntDocxRenderer,
    LegalTextModifiedError,
    MissingValueError,
    TemplateError,
)
from app.docx.ooxml import W_P, W_PPR, W_R, W_RPR, W_TBL, W_TR, iter_paragraphs, paragraph_text
from app.docx.template_tagger import EXPECTED, tag_template
from app.docx.vn_text import amount_in_words, format_percent, format_vnd, number_to_vietnamese

ROOT = Path(__file__).resolve().parents[1]
TAGGED = {"ezs": ROOT / "templates/bbnt_ezs_fb_nonop.v2.docx", "sassy": ROOT / "templates/bbnt_sassy_fb_nonop.v2.docx"}
SOURCE = {"ezs": ROOT / "templates/source/TEMPLATE BBNT EZS X BRAND - FB NON-OP.docx",
          "sassy": ROOT / "templates/source/TEMPLATE BBNT HKD SASSY X BRAND - FB NON-OP.docx"}


# =========================================================================== số -> chữ
@pytest.mark.parametrize("n, words", [
    (0, "không"), (1, "một"), (5, "năm"), (10, "mười"), (11, "mười một"), (14, "mười bốn"), (15, "mười lăm"),
    (20, "hai mươi"), (21, "hai mươi mốt"), (24, "hai mươi bốn"), (25, "hai mươi lăm"), (55, "năm mươi lăm"),
    (100, "một trăm"), (101, "một trăm linh một"), (105, "một trăm linh năm"), (110, "một trăm mười"),
    (115, "một trăm mười lăm"), (121, "một trăm hai mươi mốt"), (1000, "một nghìn"),
    (1001, "một nghìn không trăm linh một"), (1010, "một nghìn không trăm mười"),
    (1050, "một nghìn không trăm năm mươi"), (2000000, "hai triệu"),
    (1005000, "một triệu không trăm linh năm nghìn"),
    (1000005, "một triệu không trăm linh năm"),
    (15912468, "mười lăm triệu chín trăm mười hai nghìn bốn trăm sáu mươi tám"),
    (88402600, "tám mươi tám triệu bốn trăm linh hai nghìn sáu trăm"),
    (1_000_000_000, "một tỷ"),
    (1_000_000_005, "một tỷ không trăm linh năm"),
    (2_500_000_000, "hai tỷ năm trăm triệu"),
    (1_000_000_000_000, "một nghìn tỷ"),
    (1_234_567_891_011, "một nghìn hai trăm ba mươi bốn tỷ năm trăm sáu mươi bảy triệu tám trăm chín mươi mốt "
                        "nghìn không trăm mười một"),
])
def test_number_to_vietnamese(n, words):
    assert number_to_vietnamese(n) == words


def test_number_options():
    assert number_to_vietnamese(105, linh="lẻ") == "một trăm lẻ năm"
    assert number_to_vietnamese(24, four_after_tens="tư") == "hai mươi tư"
    assert number_to_vietnamese(14, four_after_tens="tư") == "mười bốn"   # 'tư' chỉ sau 'mươi'


def test_amount_in_words():
    assert amount_in_words(17185465) == "Mười bảy triệu một trăm tám mươi lăm nghìn bốn trăm sáu mươi lăm đồng"
    assert amount_in_words(Decimal("14850000.00")) == "Mười bốn triệu tám trăm năm mươi nghìn đồng"
    assert amount_in_words(0) == "Không đồng"
    with pytest.raises(ValueError, match="phần lẻ"):
        amount_in_words(Decimal("1272997.44"))
    with pytest.raises(ValueError):
        amount_in_words(-1)
    with pytest.raises(TypeError):
        number_to_vietnamese(1.0)  # type: ignore[arg-type]


@pytest.mark.parametrize("v, s", [(17185465, "17.185.465"), (0, "0"), (Decimal("1272997.44"), "1.272.997,44"),
                                  (Decimal("1272997.00"), "1.272.997"), (Decimal("-1000"), "-1.000"), (999, "999")])
def test_format_vnd(v, s):
    assert format_vnd(v) == s


@pytest.mark.parametrize("v, s", [(Decimal("0.18"), "18%"), (Decimal("0.08"), "8%"), (Decimal("0.053070"), "5,31%"),
                                  (Decimal("0.015"), "1,5%"), (Decimal("1"), "100%"), (Decimal("0"), "0%"), (None, "")])
def test_format_percent(v, s):
    assert format_percent(v) == s


# =========================================================================== helpers
def values_for(r: BbntDocxRenderer, **over: str) -> dict[str, str]:
    vals = {k: f"<{k}>" for k in r.spec.scalars}
    vals.update(over)
    return vals


def items(n: int) -> list[dict[str, str]]:
    return [{"name": f"SP {i}", "quantity": str(i), "revenue": format_vnd(1000 * i)} for i in range(1, n + 1)]


def load(content: bytes) -> docx.document.Document:
    return docx.Document(io.BytesIO(content))


def all_text(d) -> str:
    return "\n".join(paragraph_text(p) for p in d.element.body.iter(W_P))


def c14n(data: bytes) -> bytes:
    return etree.tostring(etree.fromstring(data), method="c14n")


def fingerprint(d, *, skip_item_rows: bool = True) -> list[tuple[bytes, tuple[bytes, ...]]]:
    """Định dạng của mọi paragraph (pPr + chuỗi rPr) bỏ qua text — để so template vs output."""
    out = []
    for p in d.element.body.iter(W_P):
        tr = next(p.iterancestors(W_TR), None)
        if skip_item_rows and tr is not None and re.search(r"\{\{item\.|^SP \d", "".join(
                paragraph_text(x) for x in tr.iter(W_P))):
            continue
        ppr = p.find(W_PPR)
        rprs = tuple(etree.tostring(r.find(W_RPR)) if r.find(W_RPR) is not None else b"" for r in p.iter(W_R))
        out.append((etree.tostring(ppr) if ppr is not None else b"", rprs))
    return out


# =========================================================================== real templates
@pytest.fixture(params=["ezs", "sassy"])
def kind(request):
    return request.param


def test_tagged_templates_expose_expected_tokens(kind):
    spec = BbntDocxRenderer(TAGGED[kind]).spec
    assert spec.scalars == EXPECTED[kind]
    assert spec.item_fields == {"name", "quantity", "revenue"}


@pytest.mark.parametrize("n_items", [0, 1, 3, 60])
def test_render_preserves_formatting_and_parts(kind, n_items):
    tpl = TAGGED[kind].read_bytes()
    r = BbntDocxRenderer(tpl)
    out = r.render(values_for(r), items(n_items), seed=7)
    assert out.item_rows == n_items

    # 1) Mọi part ngoài document.xml giống hệt (styles, numbering, footer, ảnh, fonts nhúng, comments...)
    zt, zo = zipfile.ZipFile(io.BytesIO(tpl)), zipfile.ZipFile(io.BytesIO(out.content))
    assert set(zt.namelist()) == set(zo.namelist())
    for name in zt.namelist():
        if name == "word/document.xml":
            continue
        a, b = zt.read(name), zo.read(name)
        if name.endswith((".xml", ".rels")):
            assert c14n(a) == c14n(b), name
        else:
            assert a == b, name

    # 2) document.xml: sectPr (margins/page) và định dạng mọi paragraph ngoài dòng sản phẩm giữ nguyên
    dt, do = load(tpl), load(out.content)
    sect = lambda d: etree.tostring(d.element.body.find("{*}sectPr"))  # noqa: E731
    assert sect(dt) == sect(do)
    assert fingerprint(dt) == fingerprint(do)

    # 3) Không còn token
    assert "{{" not in all_text(do)


def test_item_rows_cloned_with_row_and_cell_formatting(kind):
    tpl = TAGGED[kind].read_bytes()
    r = BbntDocxRenderer(tpl)
    d = load(r.render(values_for(r), items(5), seed=1).content)
    tbl = next(t for t in d.element.body.iter(W_TBL) if "TÊN SẢN PHẨM" in paragraph_text(next(t.iter(W_P))))
    rows = tbl.findall(W_TR)
    texts = [[paragraph_text(p) for p in tr.iter(W_P)] for tr in rows]
    assert texts[0][0] == "TÊN SẢN PHẨM" and texts[-1][0] == "TỔNG CỘNG"
    assert [t[:3] for t in texts[1:-1]] == [[f"SP {i}", str(i), format_vnd(1000 * i)] for i in range(1, 6)]
    proto = next(tr for tr in load(tpl).element.body.iter(W_TR) if "{{item.name}}" in "".join(
        paragraph_text(p) for p in tr.iter(W_P)))
    strip_ids = lambda x: re.sub(rb' w14:(paraId|textId)="[^"]*"', b"", x)  # noqa: E731

    def ser(el):
        return etree.tostring(el) if el is not None else None
    for tr in rows[1:-1]:
        assert ser(tr.find("{*}trPr")) == ser(proto.find("{*}trPr"))
        assert [ser(tc.find("{*}tcPr")) for tc in tr.findall("{*}tc")] == \
               [ser(tc.find("{*}tcPr")) for tc in proto.findall("{*}tc")]
        assert strip_ids(etree.tostring(tr.find("{*}tc/{*}p/{*}pPr"))) == \
               strip_ids(etree.tostring(proto.find("{*}tc/{*}p/{*}pPr")))
    para_ids = [p.get("{http://schemas.microsoft.com/office/word/2010/wordml}paraId") for p in d.element.body.iter(W_P)]
    para_ids = [x for x in para_ids if x]
    assert len(para_ids) == len(set(para_ids)), "paraId bị trùng sau khi clone dòng"


def test_legal_dates_untouched(kind):
    src = docx.Document(str(SOURCE[kind]))
    r = BbntDocxRenderer(TAGGED[kind])
    out = load(r.render(values_for(r), items(2)).content)
    legal = [m.group(0) for p in src.element.body.iter(W_P)
             for m in filler_mod.LEGAL_DATE_RE.finditer(paragraph_text(p))]
    assert any("tháng XX năm 20XX" in x for x in legal)          # ngày ký HĐ
    assert any(x.startswith("ngày\xa007") for x in legal)        # ngày lập BBNT
    text = all_text(out)
    for x in legal:
        assert x in text


def test_party_b_and_terms_identical_to_source(kind):
    """Mọi đoạn KHÔNG chứa placeholder (Bên B, điều khoản, chữ ký Bên B) giống hệt template gốc."""
    def texts(d):
        return [paragraph_text(p) for p in d.element.body.iter(W_P)]
    src = texts(docx.Document(str(SOURCE[kind])))
    tagged = texts(load(TAGGED[kind].read_bytes()))
    placeholder = re.compile(r"XX|xx|x%|…|\.\.\.|BRAND|^$|^ĐẠI DIỆN|Chức vụ|Địa chỉ|[Đđ]iện thoại|Mã số thuế|"
                             r"Tài khoản|Ngân hàng|[Đđ]ại diện|Số TK|BÊN A")
    kept = [t for t in src if not placeholder.search(t)]
    for t in kept:
        assert t in tagged, t
    assert tagged.count("BÊN B: " + ("CÔNG TY CỔ PHẦN THƯƠNG MẠI DỊCH VỤ EZS" if kind == "ezs" else
                                      "HỘ KINH DOANH SASSY")) == 1


def test_tagger_is_reproducible(kind):
    a = BbntDocxRenderer(tag_template(kind, SOURCE[kind].read_bytes())).spec
    assert a == BbntDocxRenderer(TAGGED[kind]).spec


def test_business_month_fields_use_mm_yy(kind):
    r = BbntDocxRenderer(TAGGED[kind])
    out = load(r.render(values_for(r, program_name="ORGANICA BDAY 07-09.26",
                                   bbnt_number="10.26/BBNT-ORGANICA-EZ-FB-CAMP-07-09.26"), []).content)
    text = all_text(out)
    assert "(Chương trình ORGANICA BDAY 07-09.26 )" in text
    assert "10.26/BBNT-ORGANICA-EZ-FB-CAMP-07-09.26" in text


def test_strip_comments_option():
    r = BbntDocxRenderer(TAGGED["ezs"])
    kept = r.render(values_for(r), []).content
    stripped = r.render(values_for(r), [], strip_comments=True).content
    assert b"commentReference" in zipfile.ZipFile(io.BytesIO(kept)).read("word/document.xml")
    assert b"commentReference" not in zipfile.ZipFile(io.BytesIO(stripped)).read("word/document.xml")


# =========================================================================== synthetic templates
def make_doc(build) -> bytes:
    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_token_split_across_runs_keeps_first_run_format():
    def build(d):
        p = d.add_paragraph("Số: ")
        r1 = p.add_run("{{bbnt_")
        r1.bold = True
        r2 = p.add_run("number}}")
        r2.italic = True
        p.add_run(" hết")
    r = BbntDocxRenderer(make_doc(build))
    assert r.spec.scalars == {"bbnt_number"}
    p = load(r.render({"bbnt_number": "06.26/BBNT-X-EZ-FB-CAMP-05.26"}).content).paragraphs[0]
    assert p.text == "Số: 06.26/BBNT-X-EZ-FB-CAMP-05.26 hết"
    assert p.runs[1].text == "06.26/BBNT-X-EZ-FB-CAMP-05.26" and p.runs[1].bold
    assert p.runs[2].text == ""      # run cũ giữ lại (rỗng) — không xoá phần tử


def test_values_are_not_reinterpreted_and_newlines_become_breaks():
    r = BbntDocxRenderer(make_doc(lambda d: d.add_paragraph("A={{a}} B={{b}}")))
    d = load(r.render({"a": "{{b}} <x>&", "b": "dòng 1\ndòng 2"}).content)
    xml = etree.tostring(d.paragraphs[0]._p).decode()
    assert d.paragraphs[0].text == "A={{b}} <x>& B=dòng 1\ndòng 2"
    assert "<w:br/>" in xml


def test_missing_and_non_string_values():
    r = BbntDocxRenderer(make_doc(lambda d: d.add_paragraph("{{a}} {{b}}")))
    with pytest.raises(MissingValueError) as e:
        r.render({"a": "1"})
    assert e.value.missing == ["b"]
    with pytest.raises(TypeError):
        r.render({"a": "1", "b": 2})  # type: ignore[dict-item]


def test_item_value_missing():
    def build(d):
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(1, 0).text, t.cell(1, 1).text = "TÊN", "{{item.name}}", "{{item.revenue}}"
    r = BbntDocxRenderer(make_doc(build))
    with pytest.raises(MissingValueError, match=r"item\[0\]\.revenue"):
        r.render({}, [{"name": "x"}])


def test_token_inside_legal_date_rejected():
    with pytest.raises(TemplateError, match="ngày pháp lý"):
        BbntDocxRenderer(make_doc(lambda d: d.add_paragraph("TPHCM, ngày {{d}} tháng 10 năm 2026")))


def test_item_token_outside_table_rejected():
    with pytest.raises(TemplateError, match="dòng bảng"):
        BbntDocxRenderer(make_doc(lambda d: d.add_paragraph("{{item.name}}")))


def test_two_prototype_rows_in_one_table_rejected():
    def build(d):
        t = d.add_table(rows=2, cols=1)
        t.cell(0, 0).text, t.cell(1, 0).text = "{{item.name}}", "{{item.name}}"
    with pytest.raises(TemplateError, match="1 dòng mẫu"):
        BbntDocxRenderer(make_doc(build))


def test_legal_text_guard_detects_tampering(monkeypatch):
    r = BbntDocxRenderer(make_doc(lambda d: (d.add_paragraph("ký ngày 01 tháng 02 năm 2026"),
                                             d.add_paragraph("{{a}}"))))
    original = BbntDocxRenderer._fill_paragraph

    def evil(p, values, item=None):
        original(p, values, item)
        for t in p.iter("{*}t"):
            t.text = (t.text or "").replace("01 tháng", "09 tháng")
    monkeypatch.setattr(BbntDocxRenderer, "_fill_paragraph", staticmethod(evil))
    with pytest.raises(LegalTextModifiedError):
        r.render({"a": "x"})


def test_empty_cell_token_inherits_paragraph_mark_font():
    """Template tagged: token ở ô trống mang font Times New Roman 12 của paragraph mark (không rơi về Calibri)."""
    d = load(TAGGED["ezs"].read_bytes())
    p = next(p for p in iter_paragraphs(d.element.body) if paragraph_text(p) == "{{commission_amount}}")
    run = next(r for r in p.iter(W_R) if "{{commission_amount}}" in "".join(t.text or "" for t in r.iter("{*}t")))
    xml = etree.tostring(run.find(W_RPR)).decode()
    assert "Times New Roman" in xml and 'w:sz w:val="24"' in xml and "<w:b" in xml
