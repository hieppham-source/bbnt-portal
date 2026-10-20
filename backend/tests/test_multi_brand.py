"""End-to-end: link order kiểu "1 tab riêng cho mỗi brand trong group" (vd KOSE: tab ANESSA/ELIXIR/DPROGRAM) —
Portal tạo 1 BBNT, danh sách sản phẩm tách riêng từng brand (doanh thu ước tính/NMV/tỷ lệ huỷ riêng), nhưng
commission vẫn tính trên TỔNG doanh thu mọi brand có trong link order. Dùng lại group SHISEIDO_GRP (SHISEIDO +
ANESSA, pháp nhân SASSY) có sẵn trong master mẫu — xem :class:`app.services.generator.BbntGenerator._read_brand_tabs`
và :func:`app.services.order_tabs.detect_brand_tabs`."""
from __future__ import annotations

import io
from decimal import Decimal

import docx
import pytest

from .conftest import SHISEIDO_Q3, FakeSheets, order_rows, sheet_id, vn_clock
from .test_api import Env, docx_text

ORDER_URL_MULTI = "https://docs.google.com/spreadsheets/d/1OrderShiseidoMultiTabAbCdEfGh/edit"

GROUP = {"channel": "FB", "entity": "SASSY", "group_code": "SHISEIDO_GRP",
        "period_months": ["07.26", "08.26", "09.26"], "campaign_name": "Q3", "order_sheet_url": ORDER_URL_MULTI}

# Tách đúng dữ liệu của test_preview_group_sassy (tab chung) thành 2 tab riêng theo brand — tổng NMV/GMV giữ
# nguyên để so chéo kết quả combined giữa 2 kiểu file order.
SHISEIDO_TAB = order_rows([["05/07/2026", "S001", "Đã giao", "Shiseido", "SHI-01", "Ultimune", 2, 10_000_000]])
ANESSA_TAB = order_rows([
    ["05/08/2026", "S002", "Đã giao", "Anessa", "ANS-01", "Anessa Gold", 5, 5_000_000],
    ["05/09/2026", "S003", "Trả hàng", "Anessa", "ANS-01", "Anessa Gold", 1, 1_000_000],
])


@pytest.fixture
def env(master_tables, fake_sheets: FakeSheets) -> Env:
    e = Env(master_tables, fake_sheets)
    e.clock = vn_clock(2026, 10, 5)
    return e


def put(env: Env, tabs: dict) -> None:
    env.sheets.sheets[sheet_id(ORDER_URL_MULTI)] = tabs


def test_brand_breakdown_computed_per_brand_commission_on_combined_total(env):
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA": ANESSA_TAB})
    r = env.post("/api/v1/bbnt/preview", GROUP)
    assert r.status_code == 200, r.text
    b = r.json()

    # giống hệt kết quả test_preview_group_sassy (file 1 tab chung, lọc theo cột Brand) — chứng tỏ 2 cách đọc
    # file order cho cùng 1 kết quả combined
    assert b["bbnt_number"] == "10.26/BBNT-SHISEIDO_GRP-SS-FB-CAMP-07-09.26"
    assert b["brands"] == ["SHISEIDO", "ANESSA"] and b["scope_code"] == "SHISEIDO_GRP"
    assert Decimal(b["figures"]["nmv"]) == 15_000_000
    assert Decimal(b["figures"]["commission_amount"]) == 2_250_000          # trên TỔNG, không phải riêng brand
    assert Decimal(b["figures"]["tax_deduction_amount"]) == 22_500
    assert Decimal(b["figures"]["total_payable"]) == 2_227_500
    # file "1 tab/brand" -> không có 1 tab chung nào được chọn
    assert b["order_tab"] is None and b["order_tabs"] == []

    breakdown = {x["brand_code"]: x for x in b["brand_breakdown"]}
    assert set(breakdown) == {"SHISEIDO", "ANESSA"}
    sh, an = breakdown["SHISEIDO"], breakdown["ANESSA"]
    assert sh["index"] == 1 and sh["order_tabs"] == ["SHISEIDO"]
    assert Decimal(sh["gmv_estimated"]) == 10_000_000 and Decimal(sh["nmv"]) == 10_000_000
    assert sh["cancellation_rate"] is None or Decimal(sh["cancellation_rate"]) == 0
    assert [i["product_name"] for i in sh["items"]] == ["Ultimune"]

    assert an["index"] == 2 and an["order_tabs"] == ["ANESSA"]
    # "Trả hàng" bị loại khỏi NMV nhưng vẫn tính vào GMV ước tính
    assert Decimal(an["gmv_estimated"]) == 6_000_000 and Decimal(an["nmv"]) == 5_000_000
    assert round(Decimal(an["cancellation_rate"]), 6) == round(Decimal(1) - Decimal(5) / Decimal(6), 6)
    assert [i["product_name"] for i in an["items"]] == ["Anessa Gold"]

    # tổng gmv_estimated/nmv 2 brand khớp đúng con số TỔNG ở "figures"
    assert Decimal(sh["gmv_estimated"]) + Decimal(an["gmv_estimated"]) == Decimal(b["figures"]["gmv_estimated"])
    assert Decimal(sh["nmv"]) + Decimal(an["nmv"]) == Decimal(b["figures"]["nmv"])


def test_missing_brand_tab_skipped_without_error(env):
    """1 brand trong group không có tab order nào trong link tháng đó -> bỏ qua, KHÔNG lỗi (theo quyết định
    của người dùng: "Bỏ qua, không hiện trong BBNT"). Thêm 1 brand thứ 3 (DPROGRAM) vào group để vẫn còn >= 2
    tab khớp RÕ RÀNG tên brand — nếu chỉ còn 1 tab, không cách nào phân biệt được với file 1 tab thường."""
    env.tables["BRAND"].append(["FB", "DPROGRAM", "Dprogram", "SHISEIDO_GRP", "Shiseido", "Có"])
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA": ANESSA_TAB})            # thiếu hẳn tab DPROGRAM
    r = env.post("/api/v1/bbnt/preview", GROUP, refresh="true")
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["brands"] == ["SHISEIDO", "ANESSA", "DPROGRAM"]
    assert Decimal(b["figures"]["nmv"]) == 15_000_000          # chỉ 2 brand có tab, DPROGRAM bị bỏ qua
    assert {x["brand_code"] for x in b["brand_breakdown"]} == {"SHISEIDO", "ANESSA"}
    assert b["other_brand_line_count"] == 0


def test_generate_fills_separate_brand_sections_in_docx(env):
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA": ANESSA_TAB})
    r = env.post("/api/v1/bbnt/generate", GROUP)
    assert r.status_code == 200, r.text
    text = docx_text(r.content)
    assert "{{" not in text
    # tên brand lấy từ tab BRAND ("Shiseido" / "Anessa", không phải mã brand viết hoa)
    assert "1. Shiseido" in text and "2. Anessa" in text
    assert text.index("1. Shiseido") < text.index("2. Anessa")      # heading theo đúng thứ tự brand trong group
    assert "Ultimune" in text and "Anessa Gold" in text
    # 3 bảng sản phẩm: khối Shiseido, khối Anessa, rồi bảng TỔNG CỘNG ngoài khối gộp đủ cả 2 (xem product
    # tables theo đúng thứ tự tài liệu — python-docx giữ nguyên thứ tự bảng top-level)
    d = docx.Document(io.BytesIO(r.content))
    product_tables = [t for t in d.tables if t.rows[0].cells[0].text == "TÊN SẢN PHẨM"]
    assert len(product_tables) == 3
    names = lambda t: [t.rows[i].cells[0].text for i in range(1, len(t.rows) - 1)]       # noqa: E731
    assert names(product_tables[0]) == ["Ultimune"]
    assert names(product_tables[1]) == ["Anessa Gold"]
    assert names(product_tables[2]) == ["Ultimune", "Anessa Gold"]      # TỔNG gộp cả 2 brand
    # hoa hồng vẫn tính trên TỔNG, không phải riêng brand
    assert text.count("2.250.000") >= 1


def test_brand_with_multiple_tabs_merged_into_one_brand_block(env):
    """Tái hiện đúng ca thực tế: group 2 brand, 1 brand (ANESSA) tách nhiều tab theo loại sản phẩm (như
    "BIO CORNER ORDER MIX" + "BIO CORNER ORDER FOOD" cùng 1 brand BIO CORNER) — phải gộp dòng của mọi tab
    đó vào CHUNG 1 khối/1 NMV cho brand này, không báo ORDER_TAB_AMBIGUOUS."""
    anessa_mix = order_rows([["05/08/2026", "S002", "Đã giao", "Anessa", "ANS-01", "Anessa Gold", 5, 5_000_000]])
    anessa_food = order_rows([["05/09/2026", "S003", "Trả hàng", "Anessa", "ANS-01", "Anessa Gold", 1, 1_000_000]])
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA ORDER MIX": anessa_mix, "ANESSA ORDER FOOD": anessa_food})
    r = env.post("/api/v1/bbnt/preview", GROUP)
    assert r.status_code == 200, r.text
    b = r.json()
    assert Decimal(b["figures"]["nmv"]) == 15_000_000            # giống hệt ca ANESSA 1 tab duy nhất
    assert Decimal(b["figures"]["commission_amount"]) == 2_250_000

    breakdown = {x["brand_code"]: x for x in b["brand_breakdown"]}
    assert set(breakdown) == {"SHISEIDO", "ANESSA"}
    an = breakdown["ANESSA"]
    assert set(an["order_tabs"]) == {"ANESSA ORDER MIX", "ANESSA ORDER FOOD"}
    assert Decimal(an["gmv_estimated"]) == 6_000_000 and Decimal(an["nmv"]) == 5_000_000
    # 1 brand = 1 khối duy nhất trong .docx dù gộp từ 2 tab — không tách thành 2 khối "ANESSA"
    assert [i["product_name"] for i in an["items"]] == ["Anessa Gold"]

    r2 = env.post("/api/v1/bbnt/generate", GROUP)
    assert r2.status_code == 200, r2.text
    text = docx_text(r2.content)
    assert text.count("1. Shiseido") + text.count("2. Anessa") == 2       # đúng 2 khối brand, không phải 3


def test_extra_order_sheet_urls_not_supported_with_brand_tab_file(env):
    """Link chính kiểu "1 tab/brand" + thêm ``extra_order_sheet_urls`` -> chưa hỗ trợ (gộp các tab brand vào 1
    link duy nhất thay vì dán thêm link)."""
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA": ANESSA_TAB})
    second_url = "https://docs.google.com/spreadsheets/d/1SecondShiseidoAbCdEfGhIjKlMnOp/edit"
    env.sheets.sheets[sheet_id(second_url)] = order_rows([
        ["05/07/2026", "S999", "Đã giao", "Shiseido", "SHI-02", "Lotion", 1, 1_000_000],
    ])
    r = env.post("/api/v1/bbnt/preview", {**GROUP, "extra_order_sheet_urls": [second_url]})
    assert r.status_code == 422 and r.json()["code"] == "TAB_MODE_NOT_SUPPORTED"


def test_preview_split_rejects_brand_tab_file(env):
    """File order kiểu "1 tab/brand" không hỗ trợ tách theo chương trình (``tab_mode='split'``) — mỗi brand đã
    tự tách khối riêng trong CÙNG 1 BBNT rồi."""
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA": ANESSA_TAB})
    r = env.post("/api/v1/bbnt/preview-split",
                {**GROUP, "tab_mode": "split", "tab_campaign_names": {"SHISEIDO": "A", "ANESSA": "B"}})
    assert r.status_code == 422 and r.json()["code"] == "TAB_MODE_NOT_SUPPORTED"


def test_single_tab_group_file_still_falls_back_to_old_behavior(env):
    """Group nhưng file order vẫn kiểu 1 tab chung (không phải 1 tab/brand) -> hành vi CŨ (tổng hợp theo cột
    Brand, không tách khối theo brand trong .docx) không đổi."""
    put(env, SHISEIDO_Q3)
    r = env.post("/api/v1/bbnt/preview", GROUP)
    assert r.status_code == 200, r.text
    b = r.json()
    assert Decimal(b["figures"]["nmv"]) == 15_000_000
    assert b["brand_breakdown"] == [] and b["breakdown_kind"] == "none"
    assert b["order_tab"] == "ORDER"


# ----------------------------------------------------------------------------- breakdown_mode (tách bảng
# theo brand/camp — khác detect_brand_tabs: đây là LỰA CHỌN TƯỜNG MINH, luôn gộp 1 BBNT, commission chung)
CAMP_MIX = order_rows([
    ["05/07/2026", "S101", "Đã giao", "Shiseido", "SHI-01", "Ultimune Mix", 1, 6_000_000],
    ["05/08/2026", "S102", "Đã giao", "Anessa", "ANS-01", "Anessa Mix", 2, 4_000_000],
])
CAMP_FOOD = order_rows([
    ["05/09/2026", "S103", "Đã giao", "Shiseido", "SHI-02", "Ultimune Food", 1, 4_000_000],
    ["05/09/2026", "S104", "Đã giao", "Anessa", "ANS-02", "Anessa Food", 1, 1_000_000],
])


def test_breakdown_mode_brand_uses_dedicated_tabs_when_available(env):
    """``breakdown_mode='brand'`` trên link ĐÃ tách sẵn tab theo brand -> giống hệt hành vi tự nhận diện cũ."""
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA": ANESSA_TAB})
    r = env.post("/api/v1/bbnt/preview", {**GROUP, "breakdown_mode": "brand"})
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["breakdown_kind"] == "brand"
    breakdown = {x["brand_code"]: x for x in b["brand_breakdown"]}
    assert set(breakdown) == {"SHISEIDO", "ANESSA"}
    assert breakdown["SHISEIDO"]["order_tabs"] == ["SHISEIDO"]
    assert Decimal(b["figures"]["nmv"]) == 15_000_000
    assert Decimal(b["figures"]["commission_amount"]) == 2_250_000


def test_breakdown_mode_brand_falls_back_to_brand_column_without_dedicated_tabs(env):
    """``breakdown_mode='brand'`` trên link 1 tab chung (không tách sẵn theo brand) -> nhóm lại theo cột Brand —
    KHÁC với hành vi mặc định (``test_single_tab_group_file_still_falls_back_to_old_behavior``: không tách)."""
    put(env, SHISEIDO_Q3)
    r = env.post("/api/v1/bbnt/preview", {**GROUP, "breakdown_mode": "brand"})
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["breakdown_kind"] == "brand"
    breakdown = {x["brand_code"]: x for x in b["brand_breakdown"]}
    assert set(breakdown) == {"SHISEIDO", "ANESSA"}
    assert Decimal(breakdown["SHISEIDO"]["nmv"]) == 10_000_000
    assert Decimal(breakdown["ANESSA"]["nmv"]) == 5_000_000
    assert Decimal(b["figures"]["nmv"]) == 15_000_000
    assert Decimal(b["figures"]["commission_amount"]) == 2_250_000       # vẫn trên TỔNG


def test_breakdown_mode_brand_fallback_works_across_multiple_links(env):
    put(env, SHISEIDO_Q3)
    second_url = "https://docs.google.com/spreadsheets/d/1SecondShiseidoGroupAbCdEfGh/edit"
    env.sheets.sheets[sheet_id(second_url)] = order_rows([
        ["05/09/2026", "S999", "Đã giao", "Shiseido", "SHI-03", "Ultimune Extra", 1, 2_000_000],
    ])
    r = env.post("/api/v1/bbnt/preview", {**GROUP, "breakdown_mode": "brand", "extra_order_sheet_urls": [second_url]})
    assert r.status_code == 200, r.text
    b = r.json()
    breakdown = {x["brand_code"]: x for x in b["brand_breakdown"]}
    assert Decimal(breakdown["SHISEIDO"]["nmv"]) == 12_000_000
    assert Decimal(breakdown["ANESSA"]["nmv"]) == 5_000_000
    assert Decimal(b["figures"]["nmv"]) == 17_000_000


def test_breakdown_mode_brand_fallback_generates_docx_with_separate_blocks(env):
    put(env, SHISEIDO_Q3)
    r = env.post("/api/v1/bbnt/generate", {**GROUP, "breakdown_mode": "brand"})
    assert r.status_code == 200, r.text
    text = docx_text(r.content)
    assert "1. Shiseido" in text and "2. Anessa" in text


def test_breakdown_mode_camp_splits_tables_by_camp_shared_commission(env):
    """``breakdown_mode='camp'``: mỗi tab (dù chứa nhiều brand trộn lẫn) là 1 camp riêng, commission vẫn tính
    trên TỔNG — không được dùng tên tab làm tên camp (``camp_names`` bắt buộc)."""
    put(env, {"CAMP MIX": CAMP_MIX, "CAMP FOOD": CAMP_FOOD})
    payload = {**GROUP, "breakdown_mode": "camp",
              "camp_names": [{"link_index": 0, "tab": "CAMP MIX", "name": "Mix"},
                             {"link_index": 0, "tab": "CAMP FOOD", "name": "Food"}]}
    r = env.post("/api/v1/bbnt/preview", payload)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["breakdown_kind"] == "camp"
    assert Decimal(b["figures"]["nmv"]) == 15_000_000
    assert Decimal(b["figures"]["commission_amount"]) == 2_250_000
    camps = {x["brand_name"]: x for x in b["brand_breakdown"]}
    assert set(camps) == {"Mix", "Food"}
    assert all(c["brand_code"] is None for c in camps.values())
    assert Decimal(camps["Mix"]["nmv"]) == 10_000_000 and Decimal(camps["Food"]["nmv"]) == 5_000_000
    assert camps["Mix"]["order_tabs"] == ["CAMP MIX"] and camps["Food"]["order_tabs"] == ["CAMP FOOD"]


def test_breakdown_mode_camp_requires_a_name_for_every_tab(env):
    put(env, {"CAMP MIX": CAMP_MIX, "CAMP FOOD": CAMP_FOOD})
    r = env.post("/api/v1/bbnt/preview", {**GROUP, "breakdown_mode": "camp",
                 "camp_names": [{"link_index": 0, "tab": "CAMP MIX", "name": "Mix"}]})
    assert r.status_code == 409 and r.json()["code"] == "TAB_CAMPAIGN_NAMES_REQUIRED"
    assert r.json()["details"]["tabs"] == [{"link_index": 0, "tab": "CAMP FOOD"}]


def test_breakdown_mode_camp_rejects_duplicate_names(env):
    put(env, {"CAMP MIX": CAMP_MIX, "CAMP FOOD": CAMP_FOOD})
    r = env.post("/api/v1/bbnt/preview", {**GROUP, "breakdown_mode": "camp",
                 "camp_names": [{"link_index": 0, "tab": "CAMP MIX", "name": "Same"},
                                {"link_index": 0, "tab": "CAMP FOOD", "name": "same "}]})
    assert r.status_code == 422 and r.json()["code"] == "TAB_CAMPAIGN_NAMES_DUPLICATE"


def test_breakdown_mode_rejects_combination_with_order_tab_or_tab_mode(env):
    put(env, {"SHISEIDO": SHISEIDO_TAB, "ANESSA": ANESSA_TAB})
    r = env.post("/api/v1/bbnt/preview", {**GROUP, "breakdown_mode": "brand", "order_tab": "SHISEIDO"})
    assert r.status_code == 422 and r.json()["code"] == "BREAKDOWN_MODE_NOT_SUPPORTED"
    r2 = env.post("/api/v1/bbnt/preview", {**GROUP, "breakdown_mode": "brand", "tab_mode": "combine"})
    assert r2.status_code == 422 and r2.json()["code"] == "BREAKDOWN_MODE_NOT_SUPPORTED"
